"""The Smart Cafe - Cafe Quant Lab dashboard.

Launch:   streamlit run app.py
Needs:    cafe.py (CafeSim + helpers) and agent.py (train, policy_from_Q) in the same folder.
All figures are SIMULATED - not real cafe sales or market data.

Charts use Altair (installed with Streamlit). Hover any mark for exact values; click a mark to pin
its details below the chart (only that chart reruns); the selector beside each chart pins any point
without the mouse. The Word evidence pack uses python-docx + matplotlib (see requirements.txt).
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
from typing import Callable, Dict, List, Optional
from urllib.parse import quote

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="The Smart Café · Café Quant Lab", page_icon="☕", layout="wide")

try:
    from agent import epsilon_at, floor_episode, policy_from_Q, train
    from cafe import CafeSim, evaluate, reachable_states, run_month, state_features, totals
except Exception as exc:  # clear, actionable message instead of a stack trace
    st.error("Could not import **cafe.py** / **agent.py**. Keep both files next to app.py.")
    st.exception(exc)
    st.stop()

# ================================================================ design tokens & constants
NAVY, CARD, TEAL, EMERALD = "#14243A", "#FFFDF9", "#1F9D8B", "#23856D"
AMBER, CORAL, BLUE, SLATE = "#D9A441", "#D96C5F", "#557DA8", "#627184"
EVAL_SEED, SENS_SEED, GAME_SEED = 20260, 31337, 4242
EVAL_MONTHS, GAME_DAYS, MIN_VISITS, N_EXP_SEEDS = 500, 10, 100, 3
DEFAULTS = dict(spoilage_cost=1.5, stockout_penalty=6.0, mean_demand=7, spoil_frac=0.25)
SB_DEFAULTS = {"sb_spoil_cost": 1.5, "sb_penalty": 6.0, "sb_mean": 7, "sb_frac": 0.25,
               "sb_episodes": 15000, "sb_alpha": 0.15, "sb_gamma": 0.95, "sb_seed": 2026}
HINT = "Hover for exact values · click a point to pin it below"
T = alt.Tooltip  # short alias

CSS = """<style>
.stApp{background:radial-gradient(circle at 8% 0%,rgba(217,164,65,.10),transparent 38%),
  radial-gradient(circle at 95% 10%,rgba(85,125,168,.08),transparent 35%),linear-gradient(180deg,#F6F3EC 0%,#F1ECE1 100%);}
.stApp::before{content:"";position:fixed;inset:0;pointer-events:none;z-index:0;
  background-image:linear-gradient(rgba(98,113,132,.045) 1px,transparent 1px),
  linear-gradient(90deg,rgba(98,113,132,.045) 1px,transparent 1px);background-size:28px 28px;}

/* ---------- hero ---------- */
.hero{position:relative;overflow:hidden;border-radius:20px;padding:20px 26px 18px 30px;margin:2px 0 12px;color:#fff;
  background:radial-gradient(560px 280px at 95% -15%,rgba(31,157,139,.34),transparent 62%),
             radial-gradient(460px 260px at -6% 120%,rgba(217,164,65,.24),transparent 62%),
             linear-gradient(125deg,#0d1a2b 0%,%NAVY% 46%,#1d3352 100%);
  box-shadow:0 18px 40px -20px rgba(20,36,58,.6),inset 0 1px 0 rgba(255,255,255,.06);border:1px solid rgba(255,255,255,.05);}
.hero::before{content:"";position:absolute;inset:0;pointer-events:none;
  background-image:linear-gradient(rgba(255,255,255,.04) 1px,transparent 1px),linear-gradient(90deg,rgba(255,255,255,.04) 1px,transparent 1px);
  background-size:26px 26px;-webkit-mask-image:linear-gradient(90deg,transparent 0%,#000 60%);mask-image:linear-gradient(90deg,transparent 0%,#000 60%);}
.hero::after{content:"";position:absolute;left:0;top:0;bottom:0;width:6px;background:linear-gradient(180deg,%TEAL%,%AMBER%);}
.hero-top{position:relative;z-index:1;display:flex;justify-content:space-between;align-items:center;gap:10px;flex-wrap:wrap}
.eyebrow{display:inline-flex;align-items:center;gap:8px;font-size:.66rem;letter-spacing:.14em;text-transform:uppercase;
  color:#a6e3d8;background:rgba(31,157,139,.14);border:1px solid rgba(31,157,139,.38);padding:4px 11px;border-radius:999px;font-weight:600}
.badge{display:inline-flex;align-items:center;gap:8px;font-size:.74rem;padding:4px 12px;border-radius:999px;
  background:rgba(255,255,255,.07);border:1px solid rgba(255,255,255,.14);color:#e7edf3;font-weight:600}
.badge i{width:8px;height:8px;border-radius:50%;background:var(--c);box-shadow:0 0 0 3px rgba(255,255,255,.08),0 0 10px var(--c)}
.hero-grid{position:relative;z-index:1;display:grid;grid-template-columns:minmax(0,1.6fr) minmax(0,1fr);gap:26px;align-items:center;margin-top:12px}
.hero-title{display:flex;align-items:center;gap:14px;margin:0 0 6px}
.hero-title h1{margin:0;padding:0;color:#fff;font-size:2.0rem;line-height:1.12;font-weight:800;letter-spacing:-.015em}
.hero-title h1 span{background:linear-gradient(90deg,#5fd3c1,#f0c56a);-webkit-background-clip:text;background-clip:text;color:transparent}
.cup{flex:0 0 auto;width:48px;height:48px;border-radius:14px;display:grid;place-items:center;
  background:linear-gradient(145deg,rgba(255,255,255,.12),rgba(255,255,255,.03));border:1px solid rgba(255,255,255,.14)}
.hero-lede{color:#cdd8e3;font-size:.96rem;line-height:1.55;margin:4px 0 14px;max-width:46rem}
.hero-lede b{color:#fff}
.steps{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px}
.step{background:rgba(255,255,255,.055);border:1px solid rgba(255,255,255,.09);border-radius:12px;padding:9px 12px}
.step .n{font-size:.6rem;letter-spacing:.14em;color:#f0c56a;text-transform:uppercase;font-weight:700}
.step .t{font-size:.86rem;color:#fff;font-weight:650;margin-top:2px}
.step .d{font-size:.73rem;color:#aebccb;margin-top:1px;line-height:1.35}
.hero-art svg{width:100%;height:auto;display:block}
.stats{position:relative;z-index:1;display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px;
  margin-top:16px;padding-top:14px;border-top:1px solid rgba(255,255,255,.09)}
.stat{min-width:0}
.stat .l{font-size:.62rem;letter-spacing:.11em;text-transform:uppercase;color:#93a6ba;font-weight:600}
.stat .v{font-size:1.22rem;font-weight:750;color:#fff;font-variant-numeric:tabular-nums;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;margin-top:2px}
.stat .d{font-size:.71rem;color:#aebccb;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.stat .v.pos{color:#7fe0cf}.stat .v.neg{color:#f2a49a}
/* the hero responds to its OWN width (sidebar open or closed), not just the screen width */
.hero-wrap{container-type:inline-size}
@container (max-width:900px){.hero-grid{grid-template-columns:1fr}.hero-art{display:none}
  .stats{grid-template-columns:repeat(2,minmax(0,1fr));row-gap:12px}}
@container (max-width:560px){.hero{padding:16px 16px 14px 20px}.steps{grid-template-columns:1fr}
  .hero-title h1{font-size:1.4rem}.cup{width:40px;height:40px}}
@media (max-width:720px){.steps{grid-template-columns:1fr}.stats{grid-template-columns:repeat(2,minmax(0,1fr))}}

/* ---------- assumptions ribbon ---------- */
.ribbon{display:grid;grid-template-columns:minmax(0,1.7fr) minmax(0,1.05fr) minmax(0,1.35fr);gap:10px;margin:0 0 4px}
.seg{background:%CARD%;border:1px solid #e6dfd2;border-radius:12px;padding:8px 13px;box-shadow:0 1px 3px rgba(20,36,58,.04)}
.seg .sl{font-size:.62rem;letter-spacing:.12em;text-transform:uppercase;color:%SLATE%;font-weight:700}
.seg .sv{font-size:.84rem;color:%NAVY%;margin-top:2px;line-height:1.45;font-variant-numeric:tabular-nums}
.seg .sv{display:flex;flex-wrap:wrap;column-gap:14px;row-gap:2px}
.seg .sv b{font-weight:700}
.it{white-space:nowrap}
.chg{color:#8a5a06;background:rgba(217,164,65,.20);padding:0 4px;border-radius:4px}
.ribbon-note{font-size:.72rem;color:%SLATE%;margin:2px 0 10px}
@media (max-width:900px){.ribbon{grid-template-columns:1fr}}

/* ---------- cards, pins, verdict ---------- */
.kpi{background:%CARD%;border-radius:12px;padding:10px 14px;box-shadow:0 2px 8px rgba(20,36,58,.08);
  border:1px solid #ece6da;border-top:4px solid var(--c);font-variant-numeric:tabular-nums;min-height:88px;margin-bottom:8px}
.kpi .l{font-size:.72rem;color:%SLATE%;text-transform:uppercase;letter-spacing:.05em}
.kpi .v{font-size:1.35rem;font-weight:700;color:%NAVY%;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.kpi .d{font-size:.72rem;color:%SLATE%}
.sect{font-size:.8rem;color:%SLATE%;text-transform:uppercase;letter-spacing:.06em;margin:10px 0 4px}
.pin{background:%CARD%;border:1px solid #e6dfd2;border-left:4px solid %TEAL%;border-radius:10px;
  padding:8px 12px;font-size:.9rem;color:%NAVY%;font-variant-numeric:tabular-nums;margin:2px 0 6px}
.pin.empty{border-left-color:#cfc7b8;color:%SLATE%}
.pin .k{font-size:.68rem;text-transform:uppercase;letter-spacing:.06em;color:%SLATE%;display:block;margin-bottom:2px}
.verdict{border-radius:12px;padding:12px 16px;font-size:1.05rem;font-weight:650;margin:6px 0 10px;
  border:1px solid #e6dfd2;border-left:6px solid var(--c);background:%CARD%;color:%NAVY%}
.sbrand{display:flex;align-items:center;gap:11px;padding:2px 0 10px;border-bottom:1px solid #e3dccf;margin-bottom:6px}
.sbrand .ic{width:38px;height:38px;border-radius:11px;background:linear-gradient(145deg,%NAVY%,#24406a);display:grid;place-items:center}
.sbrand .t{font-weight:800;color:%NAVY%;font-size:1.0rem;line-height:1.15}
.sbrand .s{font-size:.72rem;color:%SLATE%}
@media (max-width:820px){.kpi .v{font-size:1.1rem}}
</style>"""
for _k, _v in dict(NAVY=NAVY, CARD=CARD, TEAL=TEAL, AMBER=AMBER, SLATE=SLATE).items():
    CSS = CSS.replace(f"%{_k}%", _v)


def html(s: str, where=st) -> None:
    fn = getattr(where, "html", None)
    if fn:
        fn(s)
    else:  # very old Streamlit fallback
        where.markdown(s, unsafe_allow_html=True)


html(CSS)


def aed(x: float) -> str:
    return f"−AED {abs(x):,.2f}" if x < 0 else f"AED {x:,.2f}"


def saed(x: float) -> str:
    """Signed AED: +AED 1.00 / −AED 1.00 / AED 0.00."""
    return ("+" if round(x, 2) > 0 else "") + aed(x)


def pct(x: float) -> str:
    return "n/a" if x != x else f"{100 * x:.1f}%"


def cards(items) -> None:
    for col, (label, value, note, color) in zip(st.columns(len(items)), items):
        html(f"<div class='kpi' style='--c:{color}'><div class='l'>{label}</div>"
             f"<div class='v'>{value}</div><div class='d'>{note}</div></div>", col)


def items_of(d: dict) -> tuple:
    return tuple(sorted(d.items()))


def fingerprint(*parts) -> str:
    """Short deterministic id; used in chart keys so a pinned point never outlives its data."""
    return hashlib.md5(repr(parts).encode()).hexdigest()[:8]


def svg_uri(svg: str) -> str:
    """Streamlit strips inline <svg> from HTML blocks, so graphics are delivered as CSS data-URI backgrounds."""
    return 'url("data:image/svg+xml;charset=utf-8,' + quote(svg) + '")'


def cup_svg(size: int = 26) -> str:
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" stroke="#ffffff" stroke-width="1.7" '
            'stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">'
            '<path d="M4 10h12v4.5A5.5 5.5 0 0 1 10.5 20h-1A5.5 5.5 0 0 1 4 14.5V10z"/>'
            '<path d="M16 11.5h1.6a2.4 2.4 0 0 1 0 4.8H16"/>'
            '<path d="M7.5 3.5c-.9 1 .9 2 0 3M11 3.5c-.9 1 .9 2 0 3M14.5 3.5c-.9 1 .9 2 0 3" stroke="#f0c56a"/></svg>')


def hero_art() -> str:
    """Decorative inline SVG: daily-demand bars with a rising learned-profit curve and a target line."""
    bars = [38, 52, 30, 61, 44, 70, 35, 48, 82, 40, 56, 47, 66, 51]
    b = "".join(f'<rect x="{22 + i * 21}" y="{150 - h}" width="12" height="{h}" rx="3" fill="#D9A441" opacity=".28"/>'
                for i, h in enumerate(bars))
    pts = [(20, 128), (45, 121), (70, 114), (95, 101), (120, 97), (145, 83), (170, 77), (195, 68),
           (220, 61), (245, 57), (270, 52), (295, 49), (320, 47)]
    line = " ".join(f"{x},{y}" for x, y in pts)
    area = "M20,150 L" + " L".join(f"{x},{y}" for x, y in pts) + " L320,150 Z"
    grid = "".join(f'<line x1="12" x2="330" y1="{y}" y2="{y}" stroke="rgba(255,255,255,.06)"/>' for y in (44, 80, 115, 150))
    return (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 342 176" font-family="Helvetica, Arial, sans-serif">'
        '<defs><linearGradient id="hcL" x1="0" x2="1"><stop offset="0" stop-color="#1F9D8B"/><stop offset="1" stop-color="#7fe0cf"/></linearGradient>'
        '<linearGradient id="hcA" x1="0" x2="0" y1="0" y2="1"><stop offset="0" stop-color="#1F9D8B" stop-opacity=".30"/>'
        '<stop offset="1" stop-color="#1F9D8B" stop-opacity="0"/></linearGradient></defs>'
        '<rect x="1" y="1" width="340" height="174" rx="16" fill="rgba(255,255,255,.035)" stroke="rgba(255,255,255,.09)"/>'
        f'{grid}{b}<path d="{area}" fill="url(#hcA)"/>'
        '<line x1="20" x2="322" y1="41" y2="41" stroke="rgba(240,197,106,.75)" stroke-width="1.2" stroke-dasharray="4 4"/>'
        '<text x="320" y="34" text-anchor="end" font-size="9" fill="rgba(255,255,255,.62)">best fixed order</text>'
        f'<polyline points="{line}" fill="none" stroke="#1F9D8B" stroke-opacity=".25" stroke-width="9" stroke-linecap="round" stroke-linejoin="round"/>'
        f'<polyline points="{line}" fill="none" stroke="url(#hcL)" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"/>'
        '<circle cx="320" cy="47" r="9" fill="#7fe0cf" opacity=".22"/><circle cx="320" cy="47" r="4.2" fill="#7fe0cf"/>'
        '<text x="22" y="166" font-size="8.5" letter-spacing="1.4" fill="rgba(240,197,106,.9)">DAILY DEMAND</text>'
        '<text x="236" y="80" font-size="8.5" letter-spacing="1.4" fill="#8fe6d6">LEARNED PROFIT</text></svg>')


html("<style>.cup{background:" + svg_uri(cup_svg(28)) + " center/28px 28px no-repeat,"
     "linear-gradient(145deg,rgba(255,255,255,.12),rgba(255,255,255,.03))}"
     ".sbrand .ic{background:" + svg_uri(cup_svg(22)) + " center/22px 22px no-repeat,linear-gradient(145deg,#14243A,#24406a)}"
     ".hero-art{width:100%;aspect-ratio:342/176;background:" + svg_uri(hero_art()) + " center/contain no-repeat}</style>")


# ================================================================ You-vs-AI: ONE source of truth
# Convention used everywhere:  you_minus_ai = you_total - ai_total
#   > 0  -> "you" won   |   < 0 -> "ai" won   |   == 0 -> "tie"
# Rounding rule: every daily profit is rounded to the cent (exactly as the ledgers display it) and
# totals are sums of those whole cents. The difference and the winner are computed in integer
# cents, so the winner always agrees with the totals the user sees (cent-level tolerance:
# anything below half a cent is a tie).
def to_cents(x: float) -> int:
    return int(round(float(x) * 100))


def decide(you_cents: int, ai_cents: int) -> Dict[str, object]:
    diff = int(you_cents) - int(ai_cents)
    winner = "you" if diff > 0 else "ai" if diff < 0 else "tie"
    return dict(you_cents=int(you_cents), ai_cents=int(ai_cents), diff_cents=diff, winner=winner)


def challenge_result(you_rows: List[dict], ai_rows: List[dict]) -> Dict[str, object]:
    """Totals, signed difference and winner, computed once from the two daily ledgers."""
    out = decide(sum(to_cents(r["profit"]) for r in you_rows), sum(to_cents(r["profit"]) for r in ai_rows))
    out.update(you_total=out["you_cents"] / 100, ai_total=out["ai_cents"] / 100,
               you_minus_ai=out["diff_cents"] / 100, margin=abs(out["diff_cents"]) / 100)
    return out


@st.cache_resource
def _winner_self_check() -> List[str]:
    """Built-in arithmetic check of the winner convention (runs once per server; shows nothing if OK)."""
    cases = [(250.00, 200.00, 50.00, "you"), (235.90, 241.50, -5.60, "ai"), (200.00, 200.00, 0.00, "tie"),
             (100.01, 100.00, 0.01, "you"), (100.00, 100.01, -0.01, "ai"), (100.004, 100.00, 0.00, "tie")]
    problems = []
    for you, ai, diff, who in cases:
        r = decide(to_cents(you), to_cents(ai))
        if r["winner"] != who or r["diff_cents"] != to_cents(diff):
            problems.append(f"{you} vs {ai}: got {r}")
    return problems


_bad = _winner_self_check()
if _bad:
    st.error("Internal winner-logic check failed: " + "; ".join(_bad))


# ================================================================ interactive chart helpers
def styled(ch, title: str, subtitle=HINT):
    """Shared chart theme. Titles are kept short so they fit phone-width screens; detail goes in subtitles."""
    sub = subtitle if isinstance(subtitle, list) else [subtitle]
    return (ch.properties(title=alt.TitleParams(title, subtitle=sub))
            .configure(background=CARD, padding={"left": 6, "right": 14, "top": 6, "bottom": 4})
            .configure_axis(labelColor=SLATE, titleColor=NAVY, gridColor="#ece6da", domainColor="#bdb5a7",
                            tickColor="#bdb5a7", labelFontSize=11, titleFontSize=12, titleFontWeight=600)
            .configure_title(color=NAVY, fontSize=14, fontWeight="bold", anchor="start",
                             subtitleColor=SLATE, subtitleFontSize=11, subtitlePadding=4)
            .configure_legend(labelColor=NAVY, titleColor=NAVY, orient="bottom", labelFontSize=11,
                              titleFontSize=11, symbolStrokeWidth=2, labelLimit=320)
            .configure_view(strokeWidth=0))


def pick_param(name: str):
    """Single-point click selection keyed on the row id; double-click clears."""
    return alt.selection_point(name=name, fields=["_id"], on="click", clear="dblclick", toggle=False)


def _norm(val, lookup: dict):
    if val is None:
        return None
    if val in lookup:
        return val
    for cast in (int, str):
        try:
            if cast(val) in lookup:
                return cast(val)
        except (TypeError, ValueError):
            pass
    return None


def inspect_chart(chart, key: str, params: List[str], lookup: Dict, describe: Callable[[dict], str],
                  pick_key: Optional[str] = None, picker: Optional[dict] = None, render_pin: bool = True):
    """Render an Altair chart whose clicks write to ONE persistent pick value (st.session_state).

    * hover tooltips come from the chart itself;
    * a click on any point stores its row id; clicking empty space clears it;
    * an optional selector (picker) writes to the same pick value - the keyboard/precise fallback;
    * the pinned row is described in a box below the chart.
    """
    pick_key = pick_key or f"{key}__pick"

    def _on_select() -> None:
        state = st.session_state.get(key) or {}
        sel = state.get("selection", {}) if hasattr(state, "get") else {}
        val = None
        for name in params:
            pts = sel.get(name) or []
            if isinstance(pts, list) and pts:
                val = pts[-1].get("_id")
                break
        st.session_state[pick_key] = _norm(val, lookup)

    st.altair_chart(chart, key=key, on_select=_on_select, selection_mode=list(params), theme=None,
                    width="stretch")
    if not render_pin:
        return
    if _norm(st.session_state.get(pick_key), lookup) is None:   # stale or empty -> clear safely
        st.session_state[pick_key] = None
    left, right = st.columns([3, 1.25])
    if picker:
        with right:
            if picker["kind"] == "select":
                st.selectbox(picker["label"], [None] + list(picker["options"]), key=pick_key,
                             format_func=lambda v: "— none —" if v is None else picker["fmt"](v))
            else:
                st.number_input(picker["label"], min_value=picker["min"], max_value=picker["max"], step=1,
                                key=pick_key, placeholder="type a number")
    cur = _norm(st.session_state.get(pick_key), lookup)
    with left:
        if cur is None:
            html("<div class='pin empty'><span class='k'>Pinned point</span>Nothing pinned yet — click any "
                 "point on the chart" + (" or choose one on the right" if picker else "") + ".</div>")
        else:
            html(f"<div class='pin'><span class='k'>Pinned point</span>{describe(lookup[cur])}</div>")


@st.fragment
def chart_panel(specs: List[dict]) -> None:
    """Fragment: a click or picker change re-runs ONLY these charts, never the whole dashboard
    (no retraining, no game moves, no demand regeneration). Charts that share a pin go together."""
    for spec in specs:
        inspect_chart(**spec)


def panel(chart, key, params, lookup, describe, **kw) -> None:
    chart_panel([dict(chart=chart, key=key, params=params, lookup=lookup, describe=describe, **kw)])


# ================================================================ sidebar
for _k, _v in SB_DEFAULTS.items():
    st.session_state.setdefault(_k, _v)


def reset_defaults() -> None:
    for k, v in SB_DEFAULTS.items():
        st.session_state[k] = v


with st.sidebar:
    html("<div class='sbrand'><div class='ic'></div><div><div class='t'>Café Quant Lab</div>"
         "<div class='s'>Controls · business &amp; learning</div></div></div>")
    st.header("🏪 Café economics")
    spoil_cost = st.slider("Spoilage cost (AED / spoiled unit)", 0.0, 5.0, step=0.1, key="sb_spoil_cost")
    penalty = st.slider("Stockout penalty (AED / unmet unit)", 0.0, 12.0, step=0.5, key="sb_penalty")
    mean_d = st.slider("Typical daily demand (Poisson mean)", 3, 12, step=1, key="sb_mean")
    spoil_frac = st.slider("Share of leftovers that spoil overnight", 0.0, 1.0, step=0.05, key="sb_frac")
    st.caption("Fixed: price AED 5, order AED 1, holding AED 0.30, busy day 20% (+Poisson 5), "
               "fridge 15, start stock 5, 30 days.")
    st.header("🧠 Learning")
    episodes = st.slider("Training episodes (months)", 2000, 20000, step=1000, key="sb_episodes")
    alpha = st.slider("Learning rate α", 0.01, 0.50, step=0.01, key="sb_alpha")
    gamma = st.slider("Discount γ (value of the future)", 0.0, 0.99, step=0.01, key="sb_gamma")
    seed = int(st.number_input("Training seed", 0, 10**6, step=1, key="sb_seed"))
    st.caption("ε schedule (assignment default): start 1.0, ×0.9998 per episode, floor 0.05.")
    train_clicked = st.button("🚀 Train / retrain agent", type="primary", key="btn_train", width="stretch")
    st.button("↺ Reset to assignment defaults", key="btn_reset", on_click=reset_defaults, width="stretch")

ENV = items_of(dict(spoilage_cost=spoil_cost, stockout_penalty=penalty,
                    mean_demand=mean_d, spoil_frac=spoil_frac))
CFG = dict(env=ENV, episodes=episodes, alpha=alpha, gamma=gamma, seed=seed)
P = CafeSim(**dict(ENV)).params()


# ================================================================ cached computations
@st.cache_data(show_spinner="Training a Q-learning agent (≈10–15 s)…", max_entries=48)
def cached_train(env_items, episodes, alpha, gamma, seed, eps_decay=0.9998):
    Q, curve, diag = train(CafeSim(**dict(env_items)), episodes=episodes, alpha=alpha, gamma=gamma,
                           seed=seed, eps_decay=eps_decay, return_diagnostics=True)
    return Q, np.asarray(curve), diag


@st.cache_data(max_entries=32, show_spinner=False)
def demand_matrix(env_items, months, seed):
    return CafeSim(**dict(env_items)).sample_demands(months, seed=seed)


@st.cache_data(max_entries=96, show_spinner=False)
def eval_pi(env_items, pi, seed=EVAL_SEED):
    p = CafeSim(**dict(env_items)).params()
    return evaluate(lambda s, t: pi[s], demand_matrix(env_items, EVAL_MONTHS, seed), p)


@st.cache_data(max_entries=16, show_spinner="Evaluating every policy on the same demand months…")
def scoreboard(env_items, pi):
    """Every policy and every constant order evaluated on the SAME held-out demand months."""
    p = CafeSim(**dict(env_items)).params()
    D = demand_matrix(env_items, EVAL_MONTHS, EVAL_SEED)
    g = np.random.default_rng(EVAL_SEED + 1)  # random manager's own stream; demand stays common
    pols = {"RL agent (greedy)": lambda s, t: pi[s],
            "Cautious (order 4)": lambda s, t: 4,
            "Generous (order 10 · guide)": lambda s, t: 10,
            "Generous (order 12 · original dashboard)": lambda s, t: 12,
            "Random (uniform 0–15)": lambda s, t: int(g.integers(0, int(p["max_stock"]) + 1))}
    res = {k: evaluate(f, D, p) for k, f in pols.items()}
    const = {k: evaluate(lambda s, t, k=k: k, D, p) for k in range(int(p["max_stock"]) + 1)}
    return res, const


def policy_summary(pi, visits, reach) -> Dict[str, object]:
    """Plain-language rule summary computed from the actual policy (never a forced pattern)."""
    learned = [s for s in range(16) if s in reach and visits[s] >= MIN_VISITS]
    weak = [s for s in range(16) if s not in learned]
    if not learned:
        return dict(text="", learned=[], weak=weak, target=None)
    post = [min(s + pi[s], 15) for s in learned]
    mono = all(pi[a] >= pi[b] for a, b in zip(learned, learned[1:]))
    lo_stock, hi_stock = learned[:max(1, len(learned) // 2)], learned[max(1, len(learned) // 2):] or learned[-1:]
    lo_avg, hi_avg = np.mean([pi[s] for s in lo_stock]), np.mean([pi[s] for s in hi_stock])
    target = int(np.median(post))
    text = (f"Across well-visited stock levels {learned[0]}–{learned[-1]}, the agent orders on average **{lo_avg:.1f}** "
            f"units when stock is in the lower half ({lo_stock[0]}–{lo_stock[-1]}) and **{hi_avg:.1f}** in the upper half "
            f"({hi_stock[0]}–{hi_stock[-1]}), bringing shelves to **{min(post)}–{max(post)} units** after ordering "
            f"(median {target}). "
            + ("Orders fall as stock rises. " if lo_avg > hi_avg else "Orders do **not** fall as stock rises. ")
            + ("Orders never rise with stock, a clean order-up-to pattern." if mono else
               "The rule is not perfectly monotone: neighbouring stock levels can differ by several units. This is "
               "typical of tabular Q-learning with noisy rewards, where several orders have nearly equal value; the "
               "shelf-after-ordering line is the clearer view of the strategy."))
    return dict(text=text, learned=learned, weak=weak, target=target, post_min=min(post), post_max=max(post))


if train_clicked:
    Q, curve, diag = cached_train(ENV, episodes, alpha, gamma, seed)
    st.session_state["model"] = dict(cfg=CFG, Q=Q, curve=curve, diag=diag)
MODEL = st.session_state.get("model")
STALE = MODEL is not None and MODEL["cfg"] != CFG
FRESH = MODEL is not None and not STALE
PI = tuple(int(a) for a in policy_from_Q(MODEL["Q"])) if MODEL else None
FP = fingerprint(MODEL["cfg"]) if MODEL else "none"
SB = scoreboard(ENV, PI) if FRESH else None
BEST_K = max(SB[1], key=lambda k: SB[1][k]["mean_profit"]) if SB else None


def need_model() -> bool:
    """Guard: block RL views when no model exists or it was trained on other settings."""
    if MODEL is None:
        st.info("👈 Press **Train / retrain agent** in the sidebar to populate this section.")
        return False
    if STALE:
        st.warning("Settings changed since training, so these results are hidden to avoid showing a "
                   "policy evaluated on economics it never learned. Press **Train / retrain agent**. "
                   f"Model was trained on: {dict(MODEL['cfg']['env'])}, episodes={MODEL['cfg']['episodes']}, "
                   f"α={MODEL['cfg']['alpha']}, γ={MODEL['cfg']['gamma']}, seed={MODEL['cfg']['seed']}.")
        return False
    return True


# ================================================================ header (hero + assumptions ribbon)
cu_cost = P["price"] - P["order_cost"] + P["stockout_penalty"]                          # cost of one lost sale
co_cost = P["holding_cost"] + P["spoil_frac"] * (P["spoilage_cost"] + P["order_cost"])  # rough cost of one wasted unit
if FRESH:
    _r = SB[0]["RL agent (greedy)"]
    _best = SB[1][BEST_K]["mean_profit"]
    _gap = _r["mean_profit"] - _best
    badge = (TEAL, "Agent trained · current settings")
    stats = [("RL profit / month", aed(_r["mean_profit"]), f"± {_r['std_profit']:,.0f} across {EVAL_MONTHS} held-out months", ""),
             ("Service level", pct(_r["service"]), "units sold ÷ units demanded", ""),
             ("vs best fixed order", saed(_gap), f"best fixed = {BEST_K}/day · {aed(_best)}", "pos" if _gap > 0 else "neg" if _gap < 0 else ""),
             ("Training run", f"{MODEL['cfg']['episodes']:,} episodes", f"seed {MODEL['cfg']['seed']} · final ε {MODEL['diag']['eps'][-1]:.3f}", "")]
else:
    badge = (AMBER, "Settings changed · retrain needed") if STALE else ("#9aa8b6", "Not trained yet · press Train")
    stats = [("Unit margin", aed(P["price"] - P["order_cost"]), "price − order cost", ""),
             ("Cost of a lost sale", aed(cu_cost), "margin + stockout penalty", ""),
             ("Cost of a wasted unit", "≈ " + aed(co_cost), "holding + expected spoilage", ""),
             ("Decision space", "16 × 16", "stock states × order sizes", "")]
html(
    "<div class='hero-wrap'><div class='hero'><div class='hero-top'>"
    "<span class='eyebrow'>Reinforcement learning lab · perishable inventory</span>"
    f"<span class='badge' style='--c:{badge[0]}'><i></i>{badge[1]}</span></div>"
    "<div class='hero-grid'><div>"
    "<div class='hero-title'><div class='cup'></div><h1>The Smart Café · <span>Café Quant Lab</span></h1></div>"
    "<p class='hero-lede'>Every morning the café makes one decision: <b>how many units of a perishable item to order</b>. "
    "Too many spoil, too few lose sales and disappoint customers. A Q-learning agent learns the ordering rule from "
    "profit alone. All figures are <b>simulated</b>, not real café data.</p>"
    "<div class='steps'>"
    "<div class='step'><div class='n'>01 · Observe</div><div class='t'>Stock on hand</div><div class='d'>0–15 units carried over from yesterday</div></div>"
    "<div class='step'><div class='n'>02 · Act</div><div class='t'>Choose an order</div><div class='d'>0–15 units, capped by the 15-unit fridge</div></div>"
    "<div class='step'><div class='n'>03 · Learn</div><div class='t'>Profit updates Q</div><div class='d'>revenue minus order, holding, spoilage and stockout costs</div></div>"
    "</div></div>"
    "<div class='hero-art' role='img' aria-label='Illustration: daily demand bars and a rising learned-profit curve'></div></div>"
    "<div class='stats'>" + "".join(
        f"<div class='stat'><div class='l'>{l}</div><div class='v {cls}'>{v}</div><div class='d'>{d}</div></div>"
        for l, v, d, cls in stats) + "</div></div></div>")


def _fmt(key: str, val) -> str:
    """Highlight values that differ from the assignment defaults."""
    txt = {"spoilage_cost": f"{val:.2f}", "stockout_penalty": f"{val:.2f}", "mean_demand": f"{val:g}",
           "spoil_frac": f"{val:.0%}"}[key]
    if key in DEFAULTS and val != DEFAULTS[key]:
        return f"<b class='chg' title='Assignment default: {DEFAULTS[key]}'>{txt}</b>"
    return f"<b>{txt}</b>"


_changed = any(P[k] != v for k, v in DEFAULTS.items())


def _it(label: str, value: str) -> str:
    return f"<span class='it'>{label} {value}</span>"


html("<div class='ribbon'>"
     "<div class='seg'><div class='sl'>Unit economics · AED per unit</div><div class='sv'>"
     + _it("Price", f"<b>{P['price']:.2f}</b>") + _it("Order", f"<b>{P['order_cost']:.2f}</b>")
     + _it("Holding", f"<b>{P['holding_cost']:.2f}</b>") + _it("Spoilage", _fmt('spoilage_cost', P['spoilage_cost']))
     + _it("Stockout", _fmt('stockout_penalty', P['stockout_penalty'])) + "</div></div>"
     "<div class='seg'><div class='sl'>Customer demand</div><div class='sv'>"
     + _it("Base", f"Poisson({_fmt('mean_demand', P['mean_demand'])})") + _it("Busy days", f"<b>{P['busy_chance']:.0%}</b>")
     + _it("Busy extra", f"Poisson(<b>{P['busy_extra']:g}</b>)") + "</div></div>"
     "<div class='seg'><div class='sl'>Operations</div><div class='sv'>"
     + _it("Fridge", f"<b>{P['max_stock']}</b>") + _it("Start", f"<b>{P['start_stock']}</b>")
     + _it("Spoils overnight", _fmt('spoil_frac', P['spoil_frac'])) + _it("Month", f"<b>{P['days']}</b> days") + "</div></div></div>"
     + ("<div class='ribbon-note'>Highlighted values differ from the assignment defaults.</div>" if _changed
        else "<div class='ribbon-note'>All values are the assignment defaults.</div>"))
if STALE:
    st.warning("**Model is stale** – sidebar settings differ from the ones it was trained on. Press **Train / retrain agent**.")

tabs = st.tabs(["A · Command Centre", "B · Learning Lab", "C · Learned Strategy", "D · Ledger",
                "E · Scoreboard", "F · You vs AI", "G · Experiments", "H · Assignment",
                "I · ML Diagnostics"])

# ================================================================ A · Command Centre
with tabs[0]:
    if need_model():
        res, const = SB
        r = res["RL agent (greedy)"]
        html("<div class='sect'>Headline KPIs · RL policy on "
             f"{EVAL_MONTHS} held-out simulated months</div>")
        cards([("Avg monthly net profit", aed(r["mean_profit"]), f"± {r['std_profit']:,.0f} std across months", EMERALD),
               ("Service level", pct(r["service"]), "units sold ÷ units demanded", TEAL),
               ("Spoilage KPI", pct(r["spoilage_kpi"]), "units spoiled ÷ units ordered", AMBER),
               ("Stockout rate", pct(r["stockout_rate"]), "unmet units ÷ units demanded", CORAL)])
        html("<div class='sect'>Monthly P&amp;L (average month)</div>")
        cards([("Revenue", aed(r["mean_revenue"]), f"{r['mean_sold']:.1f} units sold", TEAL),
               ("Order cost", aed(r["mean_order_cost"]), f"{r['mean_order']:.1f} units ordered", SLATE),
               ("Holding cost", aed(r["mean_holding_cost"]), f"{r['mean_leftover']:.1f} leftover units", SLATE),
               ("Spoilage cost", aed(r["mean_spoilage_cost"]), f"{r['mean_spoiled']:.1f} units spoiled", AMBER),
               ("Stockout penalty", aed(r["mean_stockout_penalty"]), f"{r['mean_stockout']:.1f} unmet units", CORAL)])
        cards([("Profit per day", aed(r["profit_per_day"]), "monthly profit ÷ 30", EMERALD),
               ("Net margin", pct(r["net_margin"]), "net profit ÷ revenue", TEAL),
               ("Loss-making months", pct(r["loss_share"]), "share of months with profit < 0", CORAL),
               ("5th-percentile month", aed(r["p5"]), "a bad (1-in-20) month", SLATE)])

        # ---- P&L waterfall: revenue -> each cost -> net profit (average month)
        steps = [("Revenue", r["mean_revenue"]), ("Order cost", -r["mean_order_cost"]),
                 ("Holding cost", -r["mean_holding_cost"]), ("Spoilage cost", -r["mean_spoilage_cost"]),
                 ("Stockout penalty", -r["mean_stockout_penalty"])]
        wrows, run = [], 0.0
        for name, amt in steps:
            wrows.append(dict(_id=name, step=name, start=run, end=run + amt, amount=amt, running=run + amt,
                              share=abs(amt) / r["mean_revenue"] if r["mean_revenue"] else float("nan")))
            run += amt
        wrows.append(dict(_id="Net profit", step="Net profit", start=0.0, end=run, amount=run, running=run,
                          share=run / r["mean_revenue"] if r["mean_revenue"] else float("nan")))
        wdf = pd.DataFrame(wrows)
        wdf["label"] = [(f"{a:+,.0f}" if s_ not in ("Revenue", "Net profit") else f"{a:,.0f}").replace("-", "−")
                        for s_, a in zip(wdf["step"], wdf["amount"])]
        wdf["top"] = wdf[["start", "end"]].max(axis=1)
        order_ = list(wdf["step"])
        wcol = alt.Color("step:N", legend=None, scale=alt.Scale(domain=order_, range=[TEAL, SLATE, "#8b97a6", AMBER, CORAL, EMERALD]))
        wtips = [T("step:N", title="Line item"), T("amount:Q", title="Amount per average month (AED)", format="+,.2f"),
                 T("running:Q", title="Running total after this step (AED)", format=",.2f"),
                 T("share:Q", title="Share of revenue", format=".1%")]
        wx = alt.X("step:N", sort=order_, title=None, scale=alt.Scale(paddingInner=0.28),
                   axis=alt.Axis(labelAngle=0, labelLimit=90, labelExpr="datum.label == 'Net profit' ? 'Net profit' : split(datum.label, ' ')[0]"))
        wbars = alt.Chart(wdf).mark_bar(cornerRadius=3).encode(
            x=wx, y=alt.Y("start:Q", title="AED per average month"), y2="end:Q", color=wcol,
            tooltip=wtips).add_params(pick_param("pick_wf"))
        wlab = alt.Chart(wdf).mark_text(dy=-8, fontSize=11, fontWeight="bold", color=NAVY).encode(
            x=wx, y="top:Q", text="label:N", tooltip=wtips)
        lookup_wf = {r_["_id"]: r_ for r_ in wrows}
        panel(styled((wbars + wlab).properties(height=300), "From revenue to net profit",
                     [f"RL policy · average of {EVAL_MONTHS} held-out months (AED)", HINT]),
              f"ch_wf_{FP}", ["pick_wf"], lookup_wf,
              lambda w: (f"<b>{w['step']}</b>: {saed(w['amount']) if w['step'] not in ('Revenue', 'Net profit') else aed(w['amount'])} "
                         f"per average month ({pct(w['share'])} of revenue) · running total after this step <b>{aed(w['running'])}</b>"),
              picker=dict(kind="select", label="Pin a line item", options=order_, fmt=str))

        # ---- distribution of monthly profit (risk view)
        prof = np.asarray(r["profits"])
        counts, edges = np.histogram(prof, bins=20)
        hrows = [dict(_id=i, lo=float(edges[i]), hi=float(edges[i + 1]), months=int(counts[i]),
                      share=counts[i] / len(prof), rng=f"{aed(edges[i])} to {aed(edges[i + 1])}") for i in range(len(counts))]
        hdf = pd.DataFrame(hrows)
        htips = [T("rng:N", title="Monthly profit range"), T("months:Q", title="Months in range", format=",d"),
                 T("share:Q", title="Share of months", format=".1%")]
        hbars = alt.Chart(hdf).mark_bar(color=BLUE, opacity=.85, stroke=CARD, strokeWidth=1).encode(
            x=alt.X("lo:Q", title="Net profit in a month (AED)", bin="binned"), x2="hi:Q",
            y=alt.Y("months:Q", title="Number of months"), tooltip=htips).add_params(pick_param("pick_hist"))
        refs = pd.DataFrame([dict(x=float(prof.mean()), lbl=f"Mean {aed(prof.mean())}", c="Mean"),
                             dict(x=float(np.percentile(prof, 5)), lbl=f"5th percentile {aed(np.percentile(prof, 5))}", c="5th percentile")])
        hrule = alt.Chart(refs).mark_rule(strokeDash=[6, 4], size=2).encode(
            x="x:Q", color=alt.Color("c:N", title=None, scale=alt.Scale(domain=["Mean", "5th percentile"], range=[EMERALD, CORAL])),
            tooltip=[T("lbl:N", title="Reference line")])
        lookup_h = {r_["_id"]: r_ for r_ in hrows}
        panel(styled((hbars + hrule).properties(height=260), "How much a month can vary",
                     [f"RL policy · {len(prof)} held-out months", HINT]),
              f"ch_hist_{FP}", ["pick_hist"], lookup_h,
              lambda h: (f"Monthly profit <b>{h['rng']}</b> · <b>{h['months']}</b> of {len(prof)} months "
                         f"({pct(h['share'])})"),
              picker=dict(kind="select", label="Pin a profit range", options=list(lookup_h),
                          fmt=lambda i: f"{aed(lookup_h[i]['lo'])} to {aed(lookup_h[i]['hi'])}"))

        costs = {"order cost": r["mean_order_cost"], "holding": r["mean_holding_cost"],
                 "spoilage": r["mean_spoilage_cost"], "stockout penalty": r["mean_stockout_penalty"]}
        top = max(costs, key=costs.get)
        best = const[BEST_K]["mean_profit"]
        gap = r["mean_profit"] - best
        pressure = ("waste pressure (spoilage cost > stockout penalty)" if r["mean_spoilage_cost"] > r["mean_stockout_penalty"]
                    else "stockout pressure (stockout penalty ≥ spoilage cost)")
        tie = abs(gap) <= 0.01 * max(abs(best), 1)
        st.info(f"**Computed insight.** The largest cost line is **{top}** ({aed(costs[top])}/month), and the "
                f"policy currently faces more **{pressure}**. On identical demand the RL policy earned "
                f"{aed(r['mean_profit'])}/month vs {aed(best)} for the best constant order "
                f"({BEST_K}/day, found afterwards by brute force): difference {saed(gap)}"
                + (" — effectively a tie (within 1%)." if tie else "."))
        with st.expander("Methodology: how these numbers are produced"):
            st.markdown(
                f"- The greedy policy (no exploration) runs on **{EVAL_MONTHS} demand months generated with seed "
                f"{EVAL_SEED}**, separate from the training seed. Every policy in tab E sees the *same* months "
                "(common random numbers), so differences come from the policy, not luck.\n"
                "- Ratios are pooled: e.g. service = total sold ÷ total demand over all months.\n"
                "- **Spoilage KPI** (spoiled ÷ ordered) is not the overnight spoil fraction "
                "(spoiled ÷ leftover, a business parameter). Pooled spoiled ÷ leftover here = "
                f"{pct(r['spoil_of_leftover'])}.\n"
                "- The waterfall uses the same averages as the cards: revenue minus the four cost lines equals the "
                "average monthly net profit exactly.\n"
                "- Fixed parameters (price, costs) are inputs; everything else is a simulation estimate.")

# ================================================================ B · Learning Lab
with tabs[1]:
    if need_model():
        c, d = MODEL["curve"], MODEL["diag"]
        n_ep = len(c)
        xpad = max(1, round(n_ep * 0.006))   # ~5 px either side so episode 0 and the last episode are not clipped
        xdom = [-xpad, n_ep - 1 + xpad]
        w = max(10, min(200, n_ep // 20))
        roll = np.full(n_ep, np.nan)
        roll[w - 1:] = np.convolve(c, np.ones(w) / w, mode="valid")   # value at episode i = mean of i-w+1..i
        cdf = pd.DataFrame({"_id": np.arange(n_ep), "episode": np.arange(n_ep), "raw": c, "eps": d["eps"],
                            "rolling": roll})
        cdf["rolling_txt"] = [f"{v:,.2f}" if v == v else f"n/a (needs {w} episodes)" for v in roll]
        raw_lbl, roll_lbl = "Raw episode profit", f"Rolling mean ({w} episodes)"
        long = pd.concat([cdf.assign(series=raw_lbl, value=cdf["raw"]),
                          cdf.dropna(subset=["rolling"]).assign(series=roll_lbl, value=lambda x: x["rolling"])])
        lookup_curve = {int(i): row for i, row in zip(cdf["episode"], cdf.to_dict("records"))}
        tips = [T("episode:Q", title="Episode (0-based)", format=",d"),
                T("raw:Q", title="Raw episode profit (AED)", format=",.2f"),
                T("rolling_txt:N", title=f"Rolling mean, {w} ep. (AED)"),
                T("eps:Q", title="ε used in episode", format=".4f")]
        color = alt.Color("series:N", title=None, scale=alt.Scale(domain=[raw_lbl, roll_lbl], range=[BLUE, NAVY]))
        zoom = alt.selection_interval(bind="scales", encodings=["x"], name="zoom_curve")
        lines = alt.Chart(long).mark_line(clip=True).encode(
            x=alt.X("episode:Q", title="Episode (0-based index) — scroll to zoom, drag to pan",
                    scale=alt.Scale(domain=xdom, nice=False)),
            y=alt.Y("value:Q", title="Training profit (AED / month)"), color=color,
            strokeWidth=alt.condition(alt.datum.series == raw_lbl, alt.value(0.7), alt.value(2.2)),
            opacity=alt.condition(alt.datum.series == raw_lbl, alt.value(0.35), alt.value(1.0))).add_params(zoom)
        # one invisible full-height strip per episode: hover anywhere above an episode to read BOTH its raw
        # value and its rolling mean (one row per episode, so x and y can never be misaligned)
        hit = alt.Chart(cdf).mark_rule(strokeWidth=4, opacity=0, clip=True, color=NAVY).encode(
            x="episode:Q", tooltip=tips).add_params(pick_param("pick_curve"))
        shared = f"pick_curve_{FP}"

        def describe_ep(rw: dict) -> str:
            roll_txt = aed(rw["rolling"]) if rw["rolling"] == rw["rolling"] else rw["rolling_txt"]
            return (f"Episode <b>{int(rw['episode']):,}</b> · raw training profit <b>{aed(rw['raw'])}</b> · "
                    f"rolling mean ({w} ep.) <b>{roll_txt}</b> · ε <b>{rw['eps']:.4f}</b>")

        fe = floor_episode()
        eps_min = d["config"]["eps_min"]
        refs = pd.DataFrame({"y": [eps_min], "lbl": [f"Floor ε_min = {eps_min}"]})
        eps_line = alt.Chart(cdf).mark_line(color=BLUE, strokeWidth=1.8).encode(
            x=alt.X("episode:Q", title="Episode (0-based index)", scale=alt.Scale(domain=xdom, nice=False)),
            y=alt.Y("eps:Q", title="ε", scale=alt.Scale(domain=[0, 1])))
        eps_hit = alt.Chart(cdf).mark_rule(strokeWidth=4, opacity=0, color=NAVY).encode(
            x="episode:Q", tooltip=tips).add_params(pick_param("pick_eps"))
        floor_rule = alt.Chart(refs).mark_rule(color=AMBER, strokeDash=[6, 4], size=2).encode(
            y="y:Q", tooltip=[T("lbl:N", title="Reference"), T("y:Q", title="ε floor", format=".2f")])
        eps_layers = [eps_line, floor_rule]
        if fe is not None and fe < n_ep:
            vref = pd.DataFrame({"x": [fe], "lbl": [f"Floor first used at episode {fe:,}"]})
            eps_layers.append(alt.Chart(vref).mark_rule(color=SLATE, strokeDash=[2, 3], size=2).encode(
                x="x:Q", tooltip=[T("lbl:N", title="Reference"), T("x:Q", title="Episode", format=",d")]))
        layers = alt.layer(*eps_layers, eps_hit)
        # both charts share one pin, so they live in ONE fragment
        chart_panel([
            dict(chart=styled((lines + hit).properties(height=330), "Learning curve",
                              ["Judge the trend, not single episodes · scroll to zoom", HINT]),
                 key=f"ch_curve_{FP}", params=["pick_curve"], lookup=lookup_curve, describe=describe_ep,
                 pick_key=shared, render_pin=False),
            dict(chart=styled(layers.properties(height=190), "Exploration rate ε by episode",
                              "Hover for ε · click to pin the same episode as above"),
                 key=f"ch_eps_{FP}", params=["pick_eps"], lookup=lookup_curve, describe=describe_ep, pick_key=shared,
                 picker=dict(kind="number", label="Pin an exact episode", min=0, max=n_ep - 1))])
        st.caption(f"{n_ep:,} episodes are plotted. The light line is each month's raw training profit; the navy line "
                   f"is the mean of the last {w} episodes. Single episodes go up and down — that is normal noise, "
                   "not failure. The curve rises as ε falls, but a shared trend is not proof of cause: Q-values are "
                   "also improving. Training profit still contains random exploration, so it sits below the held-out "
                   "evaluation in tabs A and E.")
        n = max(1, n_ep // 10)
        early, late = float(np.mean(c[:n])), float(np.mean(c[-n:]))
        k1, k2, k3 = st.columns(3)
        with k1:
            cards([(f"First {n:,} episodes", aed(early), "avg training profit", BLUE)])
        with k2:
            cards([(f"Last {n:,} episodes", aed(late), "avg training profit", EMERALD)])
        with k3:
            cards([("Change", saed(late - early), "late − early (includes exploration)", SLATE)])
        cfg = d["config"]
        st.dataframe(pd.DataFrame({
            "Hyperparameter": ["Episodes", "Episode length", "α (learning rate)", "γ (discount)",
                               "ε start", "ε decay (per episode)", "ε floor", "Seed"],
            "Value": [f"{cfg['episodes']:,}", "30 days", cfg["alpha"], cfg["gamma"], cfg["eps_start"],
                      cfg["eps_decay"], cfg["eps_min"], cfg["seed"]],
            "What it controls": ["Months of practice", "Steps per episode",
                                 "How far each update moves Q toward the new target",
                                 "How much tomorrow's value counts today",
                                 "Probability of the random branch at the start",
                                 "Multiplier applied to ε after every episode",
                                 "Minimum exploration that never switches off",
                                 "Makes training reproducible"]}).astype(str), hide_index=True)
        weak = [s for s, v in enumerate(d["state_visits"]) if v < MIN_VISITS]
        st.markdown("**How to read this tab.** *Training reward* (above) mixes learning with exploration. "
                    "*Held-out evaluation* (tabs A/E) runs the greedy policy on separate months. *Policy behaviour* "
                    f"(tab C) is the stock → order rule itself. Stock levels {weak} were visited fewer than "
                    f"{MIN_VISITS} times, so treat their actions with caution (details in tab I).")

# ================================================================ C · Learned Strategy
with tabs[2]:
    if need_model():
        visits = MODEL["diag"]["state_visits"]
        Qtab = MODEL["Q"]
        reach = reachable_states(P)
        rows = []
        for s in range(16):
            f = state_features(s, P, PI[s])
            status = ("unreachable" if s not in reach else
                      "under-visited" if visits[s] < MIN_VISITS else "well visited")
            rows.append({"_id": s, "stock": s, "order": PI[s], "post": f["post_order_inventory"],
                         "qval": float(Qtab[s, PI[s]]), "visits": int(visits[s]), "evidence": status,
                         "nr": "n/r" if s not in reach else "", "util": f["utilisation"],
                         "cov": f["base_coverage"], "cov_e": f["expected_coverage"], "cap": f["remaining_capacity"]})
        pdf = pd.DataFrame(rows)
        pdf["row"] = "Learned policy"
        prow = alt.Y("row:N", title=None, axis=None)
        lookup_pol = {r_["_id"]: r_ for r_ in rows}
        ptips = [T("stock:O", title="Stock on hand (state, units)"), T("order:Q", title="Recommended order (units)"),
                 T("post:Q", title="Post-order inventory (units)"),
                 T("qval:Q", title="Q-value of that order (AED, discounted)", format=",.2f"),
                 T("visits:Q", title="Training visits", format=",d"), T("evidence:N", title="Evidence")]
        rect = alt.Chart(pdf).mark_rect(stroke=CARD, strokeWidth=2, cornerRadius=3).encode(
            x=alt.X("stock:O", title="Stock on hand this morning (state)", axis=alt.Axis(labelAngle=0)), y=prow,
            color=alt.Color("order:Q", title=["Recommended", "order (units)"],
                            scale=alt.Scale(scheme="cividis", domain=[0, 15]),
                            legend=alt.Legend(orient="right", gradientLength=70)),
            tooltip=ptips).add_params(pick_param("pick_pol"))
        txt = alt.Chart(pdf).mark_text(fontSize=14, fontWeight="bold", dy=-5).encode(
            x="stock:O", y=prow, text="order:Q", tooltip=ptips,
            color=alt.condition(alt.datum.order < 9, alt.value("white"), alt.value(NAVY)))
        nr = alt.Chart(pdf[pdf["nr"] != ""]).mark_text(fontSize=10, dy=12).encode(
            x="stock:O", y=prow, text="nr:N", tooltip=ptips,
            color=alt.condition(alt.datum.order < 9, alt.value("white"), alt.value(NAVY)))
        describe_pol = (lambda r_: (f"Stock on hand <b>{r_['stock']} units</b> → recommended order <b>{r_['order']} units</b> "
                                    f"(shelf after ordering: {r_['post']} units) · Q-value of that order "
                                    f"<b>{aed(r_['qval'])}</b> · {r_['visits']:,} training visits · {r_['evidence']}"))
        panel(styled((rect + txt + nr).properties(height=165), "Learned policy: order per stock level",
                     [HINT, "n/r = unreachable state, not learned"]),
              f"ch_pol_{FP}", ["pick_pol"], lookup_pol, describe_pol,
              picker=dict(kind="select", label="Pin a stock level", options=list(range(16)), fmt=lambda s: f"Stock {s}"))

        # ---- order-up-to view: what the shelf looks like after ordering
        summ = policy_summary(PI, visits, reach)
        ulong = pd.concat([pdf.assign(series="Recommended order", value=pdf["order"]),
                           pdf.assign(series="Shelf after ordering", value=pdf["post"])])
        ucol = alt.Color("series:N", title=None, legend=alt.Legend(symbolOpacity=1),
                         scale=alt.Scale(domain=["Recommended order", "Shelf after ordering"], range=[BLUE, NAVY]))
        ux = alt.X("stock:O", title="Stock on hand this morning (units)", axis=alt.Axis(labelAngle=0))
        uy = alt.Y("value:Q", title="Units", scale=alt.Scale(domain=[0, 16]))
        utips = [T("stock:O", title="Stock on hand (units)"), T("order:Q", title="Recommended order (units)"),
                 T("post:Q", title="Shelf after ordering (units)"), T("visits:Q", title="Training visits", format=",d"),
                 T("evidence:N", title="Evidence")]
        ubar = alt.Chart(ulong[ulong["series"] == "Recommended order"]).mark_bar(opacity=.45, cornerRadiusEnd=3, size=26).encode(
            x=ux, y=uy, color=ucol, tooltip=utips).add_params(pick_param("pick_upto_bar"))
        shelf = ulong[ulong["series"] == "Shelf after ordering"]
        uline = alt.Chart(shelf).mark_line(strokeWidth=2.4).encode(x=ux, y=uy, color=ucol)
        upts = alt.Chart(shelf).mark_point(filled=True, size=95).encode(
            x=ux, y=uy, color=ucol, opacity=alt.condition(alt.datum.evidence == "well visited", alt.value(1), alt.value(.35)),
            tooltip=utips).add_params(pick_param("pick_upto_pt"))
        ulist = [ubar, uline, upts]
        if summ["target"] is not None:
            tdf = pd.DataFrame({"y": [summ["target"]], "lbl": [f"Typical order-up-to level ≈ {summ['target']} units"]})
            ulist = [alt.Chart(tdf).mark_rule(color=TEAL, strokeDash=[6, 4], size=2).encode(
                         y="y:Q", tooltip=[T("lbl:N", title="Reference"), T("y:Q", title="Units")])] + ulist
        ulayers = alt.layer(*ulist)   # flat layer: nested layers render blank in this Vega-Lite version
        upto_sub = ([f"Teal dashed line: typical order-up-to level ≈ {summ['target']} units"] if summ["target"] is not None else []) \
            + ["Faded points: under-visited or unreachable states", HINT]
        panel(styled(ulayers.properties(height=290), "Order-up-to view", upto_sub),
              f"ch_upto_{FP}", ["pick_upto_bar", "pick_upto_pt"], lookup_pol, describe_pol,
              picker=dict(kind="select", label="Pin a stock level ", options=list(range(16)), fmt=lambda s: f"Stock {s}"))
        if summ["text"]:
            st.markdown("**Rule summary (generated from this policy).** " + summ["text"])
        st.dataframe(pd.DataFrame([{"Stock": r_["stock"], "Recommended order": r_["order"],
                                    "Post-order inventory": r_["post"], "Training visits": r_["visits"],
                                    "Evidence": r_["evidence"], "Utilisation (diag.)": r_["util"],
                                    "Coverage ÷ base mean (diag.)": r_["cov"],
                                    "Coverage ÷ expected demand (diag.)": r_["cov_e"],
                                    "Remaining capacity (diag.)": r_["cap"]} for r_ in rows]),
                     hide_index=True, height=600, column_config={
                         "Utilisation (diag.)": st.column_config.NumberColumn(format="%.2f"),
                         "Coverage ÷ base mean (diag.)": st.column_config.NumberColumn(format="%.2f"),
                         "Coverage ÷ expected demand (diag.)": st.column_config.NumberColumn(format="%.2f")})
        st.caption(f"Caution: stock levels {summ['weak']} are unreachable or visited fewer than {MIN_VISITS} times, so "
                   "their orders are not reliable. Never-visited rows keep all-zero Q-values, and argmax then returns 0 "
                   "by tie-breaking — that '0' is a default, not a learned decision. The Q-value in the tooltip is "
                   "the agent's own estimate of discounted future profit, not a confidence score.")
        with st.expander("Feature-engineering rationale: why the state is stock only"):
            st.markdown(
                "- **Model input:** only *stock on hand this morning* (0–15). Demand is independent day to day, "
                "so stock captures everything the order decision can act on.\n"
                "- **Diagnostic columns (diag.)** are one-to-one transformations of stock: utilisation = stock ÷ 15, "
                "coverage = stock ÷ mean demand (or ÷ expected demand "
                f"{P['mean_demand'] + P['busy_chance'] * P['busy_extra']:.1f} = mean + busy chance × 5), remaining "
                "capacity = 15 − stock. They help a human read the policy but add **no new information** to the "
                "Q-table, so they are **not model inputs**.\n"
                "- **No leakage:** today's demand is drawn *after* the order is chosen; it is never part of the state.")

# ================================================================ D · Ledger
with tabs[3]:
    pol_names = ["RL agent (greedy)", "Cautious (order 4)", "Generous (order 10 · guide)",
                 "Generous (order 12 · original dashboard)"]
    fixed = {"Cautious (order 4)": 4, "Generous (order 10 · guide)": 10, "Generous (order 12 · original dashboard)": 12}
    c1, c2 = st.columns([2, 1])
    choice = c1.selectbox("Who runs the café this month?", pol_names, key="ledger_policy")
    month = int(c2.number_input(f"Evaluation month (0–{EVAL_MONTHS - 1})", 0, EVAL_MONTHS - 1, 0, key="ledger_month"))
    if choice in fixed or need_model():
        pol = (lambda s, t: fixed[choice]) if choice in fixed else (lambda s, t: PI[s])
        rows = run_month(pol, demand_matrix(ENV, EVAL_MONTHS, EVAL_SEED)[month], P)
        cols = ["day", "start_stock", "order", "available", "demand", "sold", "stockout", "leftover", "spoiled",
                "carry", "revenue", "order_cost", "holding_cost", "spoilage_cost", "stockout_penalty", "profit"]
        names = ["Day", "Start stock", "Order", "Available", "Demand", "Sold", "Unmet (stockout)", "Leftover",
                 "Spoiled", "Carry-over", "Revenue", "Order cost", "Holding cost", "Spoilage cost",
                 "Stockout penalty", "Net profit"]
        df = pd.DataFrame(rows)[cols]
        df.columns = names
        money = {c: st.column_config.NumberColumn(format="%.2f") for c in names[10:]}
        st.dataframe(df, hide_index=True, column_config=money, height=420)
        t = totals(rows)
        recon = t["revenue"] - (t["order_cost"] + t["holding_cost"] + t["spoilage_cost"] + t["stockout_penalty"])
        cards([("Revenue", aed(t["revenue"]), f"{int(t['sold'])} units sold", TEAL),
               ("Order cost", aed(t["order_cost"]), f"{int(t['order'])} units ordered", SLATE),
               ("Holding", aed(t["holding_cost"]), f"{int(t['leftover'])} leftover units", SLATE),
               ("Spoilage", aed(t["spoilage_cost"]), f"{int(t['spoiled'])} units spoiled", AMBER),
               ("Stockout penalty", aed(t["stockout_penalty"]), f"{int(t['stockout'])} unmet units", CORAL),
               ("Net profit", aed(t["profit"]), "sum of 30 daily profits", EMERALD)])
        ok = abs(recon - t["profit"]) < 1e-6
        st.markdown(f"**Reconciliation:** revenue − order − holding − spoilage − stockout = **{aed(recon)}**; "
                    f"sum of daily net profit = **{aed(t['profit'])}** → "
                    + ("✅ reconciles exactly." if ok else "⚠️ does not reconcile."))
        st.caption(f"Accounting conventions (as specified): order cost is charged on every ordered unit, including "
                   f"{int(t['clipped_units'])} unit(s) this month discarded by the 15-unit fridge cap; holding cost "
                   "applies to all leftovers, including those that then spoil; spoiled units use Python's round() "
                   "(ties go to the even number, e.g. 2.5 → 2).")
        ldf = pd.DataFrame(rows)
        ldf["cum"] = ldf["profit"].cumsum()
        ldf["sign"] = np.where(ldf["profit"] >= 0, "Profit day", "Loss day")
        ldf["_id"] = ldf["day"]
        lookup_led = {int(r_["_id"]): r_ for r_ in ldf.to_dict("records")}
        ltips = [T("day:O", title="Day"), T("start_stock:Q", title="Start stock (units)"),
                 T("order:Q", title="Order (units)"), T("demand:Q", title="Demand (units)"),
                 T("sold:Q", title="Sold (units)"), T("stockout:Q", title="Unmet (units)"),
                 T("spoiled:Q", title="Spoiled (units)"), T("carry:Q", title="Carry-over (units)"),
                 T("profit:Q", title="Daily net profit (AED)", format=",.2f"),
                 T("cum:Q", title="Cumulative profit (AED)", format=",.2f")]
        bars = alt.Chart(ldf).mark_bar(opacity=0.6, cornerRadiusEnd=2).encode(
            x=alt.X("day:O", title="Day", axis=alt.Axis(labelAngle=0)),
            y=alt.Y("profit:Q", title="Daily net profit (AED) — bars", axis=alt.Axis(orient="left")),
            color=alt.Color("sign:N", title=None, scale=alt.Scale(domain=["Profit day", "Loss day"], range=[EMERALD, CORAL])),
            tooltip=ltips).add_params(pick_param("pick_led_bar"))
        cumline = alt.Chart(ldf).mark_line(color=NAVY, strokeWidth=2).encode(
            x="day:O", y=alt.Y("cum:Q", title="Cumulative profit (AED) — navy line", axis=alt.Axis(orient="right")))
        cumpts = alt.Chart(ldf).mark_point(color=NAVY, filled=True, size=55).encode(
            x="day:O", y="cum:Q", tooltip=ltips).add_params(pick_param("pick_led_pt"))
        zero = alt.Chart(pd.DataFrame({"y": [0]})).mark_rule(color=SLATE, size=1).encode(y="y:Q")
        # the cumulative line is drawn UNDER the bars so it never intercepts the pointer; bar and point of the
        # same day carry the same row, so whichever is clicked pins the same day
        panel(styled(alt.layer(cumline + cumpts, bars + zero).resolve_scale(y="independent").properties(height=300),
                     "Daily and cumulative profit", [f"{choice} · month {month}", HINT]),
              f"ch_led_{FP}_{fingerprint(choice, month, ENV)}", ["pick_led_bar", "pick_led_pt"], lookup_led,
              lambda r_: (f"Day <b>{r_['day']}</b> · start {r_['start_stock']} + order {r_['order']} → demand "
                          f"{r_['demand']}, sold {r_['sold']}, unmet {r_['stockout']}, spoiled {r_['spoiled']}, carry "
                          f"{r_['carry']} · daily net profit <b>{aed(r_['profit'])}</b> · cumulative "
                          f"<b>{aed(r_['cum'])}</b>"),
              picker=dict(kind="select", label="Pin a day", options=list(range(1, 31)), fmt=lambda d_: f"Day {d_}"))
        buf = io.StringIO()
        wr = csv.DictWriter(buf, fieldnames=cols + ["clipped_units"], extrasaction="ignore")
        wr.writeheader()
        wr.writerows(rows)
        st.download_button("⬇️ Download this ledger (CSV)", buf.getvalue().encode("utf-8"),
                           f"ledger_{'rl' if choice not in fixed else fixed[choice]}_month{month}.csv",
                           "text/csv", key="dl_ledger")


# ================================================================ You-vs-AI state (shared by tabs E, F, H)
def new_game(rnd: int) -> None:
    """Start (or replay) a challenge: snapshot the business parameters and fix the hidden demand."""
    st.session_state["game"] = dict(env=ENV, params=dict(P), round=rnd, rows=[], stock=int(P["start_stock"]),
                                    demands=CafeSim(**dict(ENV)).sample_demands(1, GAME_DAYS, GAME_SEED + rnd)[0].tolist(),
                                    celebrated=False)


def play_day() -> None:
    gm = st.session_state["game"]
    if len(gm["rows"]) < GAME_DAYS:
        t = len(gm["rows"])
        row = run_month(lambda s, _t: int(st.session_state["game_order"]), [gm["demands"][t]], gm["params"], gm["stock"])[0]
        row["day"] = t + 1
        gm["rows"].append(row)
        gm["stock"] = int(row["carry"])


GAME_RESET_MSG = None
_g = st.session_state.get("game")
if _g is None or _g["env"] != ENV:          # parameters changed -> the old challenge is no longer comparable
    if _g is not None:
        GAME_RESET_MSG = ("Business settings changed, so the challenge was restarted — both players must face "
                          "the same economics.")
    new_game(0 if _g is None else _g["round"] + 1)
GAME = st.session_state["game"]


def game_result() -> Optional[Dict[str, object]]:
    """AI replays the SAME demand from a FRESH start with the SAME parameter snapshot; one result object."""
    if len(GAME["rows"]) < GAME_DAYS or not FRESH:
        return None
    ai_rows = run_month(lambda s, t: PI[s], GAME["demands"], GAME["params"])   # start stock from snapshot
    res_ = challenge_result(GAME["rows"], ai_rows)
    res_["ai_rows"] = ai_rows
    return res_


def challenge_board(R: Dict[str, object]) -> List[dict]:
    """10-day scoreboard: every manager plays the SAME 10 demand days from the SAME start stock.
    'You' and 'RL agent' reuse the exact ledgers behind the verdict, so all numbers agree."""
    gp, dem = GAME["params"], GAME["demands"]
    rng = np.random.default_rng(GAME_SEED + 7919 + GAME["round"])          # fixed per challenge
    rand_orders = [int(rng.integers(0, int(gp["max_stock"]) + 1)) for _ in range(GAME_DAYS)]
    entries = [("You", "You", GAME["rows"]), ("RL agent (greedy)", "RL agent", R["ai_rows"])]
    for name, k in [("Cautious (order 4)", 4), ("Generous (order 10 · guide)", 10),
                    ("Generous (order 12 · original dashboard)", 12)]:
        entries.append((name, "Guessing manager", run_month(lambda s, t, k=k: k, dem, gp)))
    entries.append(("Random (uniform 0–15)", "Guessing manager", run_month(lambda s, t: rand_orders[t], dem, gp)))
    entries.append((f"Best fixed order ({BEST_K}/day · diagnostic)", "Diagnostic benchmark",
                    run_month(lambda s, t: BEST_K, dem, gp)))
    out = []
    for name, kind, rws in entries:
        cents = sum(to_cents(x["profit"]) for x in rws)
        out.append(dict(_id=name, name=name, kind=kind if (kind in ("You", "RL agent") or cents >= 0) else "Loss",
                        total=cents / 100, sold=sum(x["sold"] for x in rws), unmet=sum(x["stockout"] for x in rws),
                        spoiled=sum(x["spoiled"] for x in rws), ordered=sum(x["order"] for x in rws)))
    out.sort(key=lambda x: -x["total"])
    for x in out:
        x["rank"] = 1 + sum(1 for y in out if y["total"] > x["total"])   # ties share a rank
    return out


GAME_RES = game_result()
BOARD = challenge_board(GAME_RES) if GAME_RES is not None else None

# ================================================================ E · Scoreboard
with tabs[4]:
    if need_model():
        res, const = SB
        rl = res["RL agent (greedy)"]["mean_profit"]
        tbl = [{"Policy": k, "Profit / month": v["mean_profit"], "Std": v["std_profit"],
                "IQR": v["q75"] - v["q25"], "5th pct month": v["p5"], "Loss months": v["loss_share"],
                "Service": v["service"], "Stockout rate": v["stockout_rate"], "Spoilage KPI": v["spoilage_kpi"],
                "Δ vs RL (AED)": v["mean_profit"] - rl} for k, v in res.items()]
        v = const[BEST_K]   # same held-out months as every other row
        tbl.append({"Policy": f"Best constant order = {BEST_K} (diagnostic benchmark)",
                    "Profit / month": v["mean_profit"], "Std": v["std_profit"], "IQR": v["q75"] - v["q25"],
                    "5th pct month": v["p5"], "Loss months": v["loss_share"], "Service": v["service"],
                    "Stockout rate": v["stockout_rate"], "Spoilage KPI": v["spoilage_kpi"],
                    "Δ vs RL (AED)": v["mean_profit"] - rl})
        tbl.sort(key=lambda r_: -r_["Profit / month"])
        st.dataframe(pd.DataFrame(tbl), hide_index=True, column_config={
            **{c: st.column_config.NumberColumn(format="%.2f") for c in
               ["Profit / month", "Std", "IQR", "5th pct month", "Δ vs RL (AED)"]},
            **{c: st.column_config.NumberColumn(format="%.3f") for c in
               ["Loss months", "Service", "Stockout rate", "Spoilage KPI"]}})
        sdf = pd.DataFrame(tbl)
        sdf["_id"] = sdf["Policy"]
        sdf["kind"] = ["RL agent" if p_.startswith("RL") else "Diagnostic benchmark" if p_.startswith("Best")
                       else ("Guessing manager" if val >= 0 else "Guessing manager (loss)")
                       for p_, val in zip(sdf["Policy"], sdf["Profit / month"])]
        sdf["label"] = [aed(x) for x in sdf["Profit / month"]]
        sdf = sdf.rename(columns={"Profit / month": "profit", "Std": "std", "Service": "service",
                                  "Stockout rate": "stockout", "Spoilage KPI": "spoil", "Δ vs RL (AED)": "delta",
                                  "Loss months": "loss"})
        lookup_sb = {r_["_id"]: r_ for r_ in sdf.to_dict("records")}
        lo, hi = min(sdf["profit"].min(), 0), max(sdf["profit"].max(), 0)
        span = hi - lo or 1
        stips = [T("Policy:N", title="Manager / policy"), T("profit:Q", title="Avg monthly profit (AED)", format=",.2f"),
                 T("std:Q", title="Std across months (AED)", format=",.2f"), T("service:Q", title="Service level", format=".1%"),
                 T("stockout:Q", title="Stockout rate", format=".1%"), T("spoil:Q", title="Spoilage KPI", format=".1%"),
                 T("loss:Q", title="Loss-making months", format=".1%"), T("delta:Q", title="Δ vs RL agent (AED)", format="+,.2f")]
        ybar = alt.Y("Policy:N", sort=list(sdf["Policy"]), title=None, axis=alt.Axis(labelLimit=320))
        bars = alt.Chart(sdf).mark_bar(cornerRadiusEnd=3, height=22).encode(
            x=alt.X("profit:Q", title="Average monthly profit (AED)", scale=alt.Scale(domain=[lo - 0.22 * span, hi + 0.22 * span])),
            y=ybar,
            color=alt.Color("kind:N", title=None, scale=alt.Scale(
                domain=["RL agent", "Diagnostic benchmark", "Guessing manager", "Guessing manager (loss)"],
                range=[EMERALD, BLUE, SLATE, CORAL])),
            tooltip=stips).add_params(pick_param("pick_sb"))
        lab_pos = alt.Chart(sdf[sdf["profit"] >= 0]).mark_text(align="left", dx=5, fontSize=11, color=NAVY).encode(
            x="profit:Q", y=ybar, text="label:N", tooltip=stips)
        lab_neg = alt.Chart(sdf[sdf["profit"] < 0]).mark_text(align="right", dx=-5, fontSize=11, color=NAVY).encode(
            x="profit:Q", y=ybar, text="label:N", tooltip=stips)
        zero = alt.Chart(pd.DataFrame({"x": [0]})).mark_rule(color=SLATE).encode(x="x:Q")
        panel(styled((bars + lab_pos + lab_neg + zero).properties(height=280), "Monthly scoreboard",
                     [f"Every manager plays the same {EVAL_MONTHS} held-out demand months", HINT]),
              f"ch_sb_{FP}", ["pick_sb"], lookup_sb,
              lambda r_: (f"<b>{r_['Policy']}</b> · avg monthly profit <b>{aed(r_['profit'])}</b> (std "
                          f"{aed(r_['std'])}) · service {pct(r_['service'])} · stockout {pct(r_['stockout'])} · "
                          f"spoilage KPI {pct(r_['spoil'])} · Δ vs RL {saed(r_['delta'])}"),
              picker=dict(kind="select", label="Pin a manager", options=list(sdf["_id"]), fmt=str))

        # ---- constant-order diagnostic (each point evaluated on the same held-out months)
        kdf = pd.DataFrame([{"_id": f"order:{k}", "order": k, "profit": v["mean_profit"], "std": v["std_profit"],
                             "months": v["months"], "service": v["service"], "stockout": v["stockout_rate"],
                             "spoil": v["spoilage_kpi"], "series": "Constant daily order"} for k, v in const.items()])
        rldf = pd.DataFrame([{"_id": "rl", "profit": rl, "months": EVAL_MONTHS, "series": "RL agent (greedy)",
                              "std": res["RL agent (greedy)"]["std_profit"], "label": f"RL agent (greedy): {aed(rl)}"}])
        lookup_k = {r_["_id"]: r_ for r_ in kdf.to_dict("records")}
        lookup_k["rl"] = rldf.to_dict("records")[0]
        ktips = [T("order:Q", title="Constant daily order (units/day)"),
                 T("profit:Q", title="Avg monthly profit (AED)", format=",.2f"),
                 T("std:Q", title="Std across months (AED)", format=",.2f"),
                 T("months:Q", title="Evaluation months", format=",d"),
                 T("service:Q", title="Service level", format=".1%"), T("stockout:Q", title="Stockout rate", format=".1%")]
        kline = alt.Chart(kdf).mark_line(color=BLUE, strokeWidth=2).encode(
            x=alt.X("order:Q", title="Constant order every day (units)", scale=alt.Scale(domain=[-0.5, 15.5]),
                    axis=alt.Axis(values=list(range(16)), format="d")),
            y=alt.Y("profit:Q", title="Avg monthly profit (AED)"))
        kpts = alt.Chart(kdf).mark_point(color=BLUE, shape="square", filled=True, size=110).encode(
            x="order:Q", y="profit:Q", tooltip=ktips).add_params(pick_param("pick_const"))
        rlrule = alt.Chart(rldf).mark_rule(color=EMERALD, strokeDash=[7, 4], size=3).encode(
            y="profit:Q", tooltip=[T("series:N", title="Reference line"),
                                   T("profit:Q", title="RL avg monthly profit (AED)", format=",.2f"),
                                   T("std:Q", title="Std across months (AED)", format=",.2f"),
                                   T("months:Q", title="Evaluation months", format=",d")]).add_params(pick_param("pick_rl"))
        rltext = alt.Chart(rldf).mark_text(align="left", dx=4, dy=-8, color=EMERALD, fontWeight="bold").encode(
            x=alt.value(4), y="profit:Q", text="label:N")
        zero = alt.Chart(pd.DataFrame({"y": [0]})).mark_rule(color=SLATE, size=1).encode(y="y:Q")

        def describe_k(r_: dict) -> str:
            if r_["_id"] == "rl":
                return (f"<b>RL agent (greedy)</b> reference line · avg monthly profit <b>{aed(r_['profit'])}</b> "
                        f"(std {aed(r_['std'])}) over {r_['months']:,} evaluation months")
            return (f"Constant daily order: <b>{r_['order']} units</b> · avg monthly profit <b>{aed(r_['profit'])}</b> "
                    f"(std {aed(r_['std'])}, {r_['months']:,} evaluation months) · service {pct(r_['service'])} · "
                    f"stockout {pct(r_['stockout'])}")

        panel(styled((zero + kline + rlrule + rltext + kpts).properties(height=300), "Profit of every constant order",
                     ["Diagnostic only — the agent never saw these numbers", "Green dashed line = RL agent", HINT]),
              f"ch_const_{FP}", ["pick_const", "pick_rl"], lookup_k, describe_k,
              picker=dict(kind="select", label="Pin a constant order", options=list(lookup_k),
                          fmt=lambda i: "RL agent line" if i == "rl" else f"Order {i.split(':')[1]}/day"))
        st.caption("Each point is a brute-force evaluation run on the same held-out months, done after training "
                   "purely for comparison; the agent never saw these numbers.")
        st.markdown(f"**Why under-ordering hurts.** One unit short costs about **{aed(cu_cost)}** (lost margin "
                    f"AED {P['price'] - P['order_cost']:.2f} + penalty AED {P['stockout_penalty']:.2f}). One unit too many "
                    f"costs about **{aed(co_cost)}** (holding + the expected share that spoils). With shortages ~"
                    f"{cu_cost / max(co_cost, 1e-9):.0f}× dearer, the newsvendor critical ratio is ≈ "
                    f"{cu_cost / (cu_cost + co_cost):.2f}, which points to stocking above average demand. This is a rough "
                    "single-day guide; carry-over and the fridge cap make the real optimum a simulation question.")
        st.markdown("#### 10-day challenge scoreboard (includes you)")
        if BOARD is not None:
            who = {"you": f"you won by {aed(GAME_RES['margin'])}", "ai": f"the AI won by {aed(GAME_RES['margin'])}",
                   "tie": "a tie"}[GAME_RES["winner"]]
            st.caption(f"Every manager played your exact 10 demand days from {GAME['params']['start_stock']} units. "
                       f"You vs AI: **{who}**. Ten days is a small sample, so use the monthly scoreboard above for the "
                       "long-run comparison. Interactive chart in tab F.")
            st.dataframe(pd.DataFrame([{"Rank": b["rank"], "Manager": b["name"], "10-day profit (AED)": b["total"],
                                        "Units sold": b["sold"], "Unmet": b["unmet"], "Spoiled": b["spoiled"]} for b in BOARD]),
                         hide_index=True, column_config={"10-day profit (AED)": st.column_config.NumberColumn(format="%.2f")})
        else:
            st.info("Play the 10-day challenge in tab F — your result is then ranked here against every manager on "
                    "the same 10 days.")

# ================================================================ F · You vs AI
with tabs[5]:
    st.markdown("You run the café for **10 days**. Each morning you see your stock and choose an order; demand "
                "is revealed only after you open the shop. Then the AI — and every other manager — plays the **same "
                "10 demand days** from the same starting stock with the same economics (snapshotted when the "
                "challenge started).")
    if GAME_RESET_MSG:
        st.warning(GAME_RESET_MSG)
    day = len(GAME["rows"])
    you_so_far = sum(to_cents(r_["profit"]) for r_ in GAME["rows"]) / 100
    if day < GAME_DAYS:
        cards([("Day", f"{day + 1} of {GAME_DAYS}", f"challenge #{GAME['round'] + 1}", BLUE),
               ("Stock this morning", f"{GAME['stock']} units", "carried over from yesterday", TEAL),
               ("Your profit so far", aed(you_so_far), "sum of your ledger below", SLATE)])
        st.number_input("Your order this morning (0–15 units)", 0, 15, 8, key="game_order")
    elif FRESH:
        st.success("All 10 days played — results below.")
    b1, b2, b3 = st.columns(3)
    b1.button("▶️ Open the shop", on_click=play_day, disabled=day >= GAME_DAYS, key="game_submit", type="primary")
    b2.button("🔁 Replay the same 10 days", on_click=new_game, args=(GAME["round"],), key="game_replay")
    b3.button("🎲 New 10-day challenge", on_click=new_game, args=(GAME["round"] + 1,), key="game_new")
    if GAME["rows"]:
        st.markdown("**Your days so far**")
        st.dataframe(pd.DataFrame(GAME["rows"])[["day", "start_stock", "order", "demand", "sold", "stockout",
                                                  "spoiled", "carry", "profit"]].rename(columns={
            "day": "Day", "start_stock": "Start", "order": "Order", "demand": "Demand", "sold": "Sold",
            "stockout": "Unmet", "spoiled": "Spoiled", "carry": "Carry", "profit": "Profit (AED)"}),
            hide_index=True, column_config={"Profit (AED)": st.column_config.NumberColumn(format="%.2f")})
    if day >= GAME_DAYS and need_model():
        R = GAME_RES
        ai_rows = R["ai_rows"]
        # --- verdict: every element below reads the same result object R
        color_you = EMERALD if R["winner"] == "you" else CORAL if R["winner"] == "ai" else SLATE
        color_ai = EMERALD if R["winner"] == "ai" else CORAL if R["winner"] == "you" else SLATE
        if R["winner"] == "you":
            verdict, short, vcol = f"🏆 You won by {aed(R['margin'])}", f"You won by {aed(R['margin'])}", EMERALD
        elif R["winner"] == "ai":
            verdict, short, vcol = f"🤖 AI won by {aed(R['margin'])}", f"AI won by {aed(R['margin'])}", CORAL
        else:
            verdict, short, vcol = "🤝 Draw — identical 10-day profit (to the cent)", "Draw (identical to the cent)", SLATE
        html(f"<div class='verdict' style='--c:{vcol}'>{verdict}</div>")
        cards([("Your 10-day profit", aed(R["you_total"]), "sum of your daily ledger", color_you),
               ("AI 10-day profit", aed(R["ai_total"]), "same demand, same start", color_ai),
               ("YOU − AI", saed(R["you_minus_ai"]), "Positive = you won; negative = AI won; zero = tie.", vcol)])
        if R["winner"] == "you" and not GAME.get("celebrated"):
            GAME["celebrated"] = True
            st.balloons()
        st.markdown("**Day by day: you and the AI on identical demand**")
        st.dataframe(pd.DataFrame([{"Day": a["day"], "Demand": a["demand"], "Your start": y["start_stock"],
                                    "Your order": y["order"], "Your profit": y["profit"], "AI start": a["start_stock"],
                                    "AI order": a["order"], "AI profit": a["profit"]}
                                   for a, y in zip(ai_rows, GAME["rows"])]),
                     hide_index=True, column_config={c: st.column_config.NumberColumn(format="%.2f")
                                                     for c in ["Your profit", "AI profit"]})
        grows = []
        for who, rws in (("You", GAME["rows"]), ("AI", ai_rows)):
            cum = 0
            for r_ in rws:
                cum += to_cents(r_["profit"])
                grows.append({"_id": f"{who}:{r_['day']}", "player": who, "day": r_["day"], "start": r_["start_stock"],
                              "order": r_["order"], "demand": r_["demand"], "sold": r_["sold"], "unmet": r_["stockout"],
                              "spoiled": r_["spoiled"], "daily": to_cents(r_["profit"]) / 100, "cum": cum / 100})
        gdf = pd.DataFrame(grows)
        # small sideways offset so equal You/AI values (e.g. a draw) never hide each other's points
        gdf["xpos"] = gdf["day"] + np.where(gdf["player"] == "You", -0.12, 0.12)
        lookup_g = {r_["_id"]: r_ for r_ in grows}
        gtips = [T("player:N", title="Player"), T("day:O", title="Day"), T("start:Q", title="Start stock (units)"),
                 T("order:Q", title="Order (units)"), T("demand:Q", title="Demand (units)"), T("sold:Q", title="Sold (units)"),
                 T("unmet:Q", title="Unmet (units)"), T("daily:Q", title="Daily profit (AED)", format=",.2f"),
                 T("cum:Q", title="Cumulative profit (AED)", format=",.2f")]
        gcol = alt.Color("player:N", title=None, scale=alt.Scale(domain=["You", "AI"], range=[AMBER, EMERALD]))
        gx = alt.X("xpos:Q", title="Day", scale=alt.Scale(domain=[0.5, GAME_DAYS + 0.5], nice=False),
                   axis=alt.Axis(values=list(range(1, GAME_DAYS + 1)), format="d", grid=False))
        gline = alt.Chart(gdf).mark_line(strokeWidth=2.2).encode(
            x=gx, y=alt.Y("cum:Q", title="Cumulative profit (AED)"),
            color=gcol, strokeDash=alt.StrokeDash("player:N", legend=None,
                                                  scale=alt.Scale(domain=["You", "AI"], range=[[1, 0], [6, 3]])))
        gpts = alt.Chart(gdf).mark_point(filled=True, size=80).encode(
            x=gx, y="cum:Q", color=gcol, shape=alt.Shape("player:N", legend=None,
                                                         scale=alt.Scale(domain=["You", "AI"], range=["circle", "square"])),
            tooltip=gtips).add_params(pick_param("pick_game"))
        panel(styled((gline + gpts).properties(height=290), "You vs AI: cumulative profit", [short, HINT]),
              f"ch_game_{GAME['round']}_{FP}", ["pick_game"], lookup_g,
              lambda r_: (f"<b>{r_['player']}</b> · day <b>{r_['day']}</b> · start {r_['start']} + order "
                          f"{r_['order']} vs demand {r_['demand']} → sold {r_['sold']}, unmet {r_['unmet']}, spoiled "
                          f"{r_['spoiled']} · daily profit <b>{aed(r_['daily'])}</b> · cumulative "
                          f"<b>{aed(r_['cum'])}</b>"),
              picker=dict(kind="select", label="Pin a player-day", options=list(lookup_g),
                          fmt=lambda i: f"{i.split(':')[0]} · day {i.split(':')[1]}"))

        # ---- 10-day scoreboard including the human
        bdf = pd.DataFrame(BOARD)
        bdf["label"] = [aed(x) for x in bdf["total"]]
        lookup_b = {b["_id"]: b for b in BOARD}
        blo, bhi = min(bdf["total"].min(), 0), max(bdf["total"].max(), 0)
        bspan = bhi - blo or 1
        btips = [T("name:N", title="Manager"), T("rank:Q", title="Rank (1 = best)"),
                 T("total:Q", title="10-day profit (AED)", format=",.2f"), T("ordered:Q", title="Units ordered"),
                 T("sold:Q", title="Units sold"), T("unmet:Q", title="Unmet demand (units)"), T("spoiled:Q", title="Spoiled (units)")]
        by = alt.Y("name:N", sort=list(bdf["name"]), title=None, axis=alt.Axis(labelLimit=320))
        bbar = alt.Chart(bdf).mark_bar(cornerRadiusEnd=3, height=20).encode(
            x=alt.X("total:Q", title="Profit over the 10 challenge days (AED)",
                    scale=alt.Scale(domain=[blo - 0.25 * bspan, bhi + 0.25 * bspan])),
            y=by, color=alt.Color("kind:N", title=None, scale=alt.Scale(
                domain=["You", "RL agent", "Guessing manager", "Diagnostic benchmark", "Loss"],
                range=[AMBER, EMERALD, SLATE, BLUE, CORAL])),
            tooltip=btips).add_params(pick_param("pick_board"))
        bpos = alt.Chart(bdf[bdf["total"] >= 0]).mark_text(align="left", dx=5, color=NAVY, fontSize=11).encode(
            x="total:Q", y=by, text="label:N", tooltip=btips)
        bneg = alt.Chart(bdf[bdf["total"] < 0]).mark_text(align="right", dx=-5, color=NAVY, fontSize=11).encode(
            x="total:Q", y=by, text="label:N", tooltip=btips)
        bzero = alt.Chart(pd.DataFrame({"x": [0]})).mark_rule(color=SLATE).encode(x="x:Q")
        you_rank = next(b["rank"] for b in BOARD if b["name"] == "You")
        panel(styled((bbar + bpos + bneg + bzero).properties(height=280), "10-day scoreboard",
                     [f"Same 10 days for every manager · you ranked {you_rank} of {len(BOARD)}", HINT]),
              f"ch_board_{GAME['round']}_{FP}", ["pick_board"], lookup_b,
              lambda b: (f"<b>{b['name']}</b> · rank <b>{b['rank']}</b> of {len(BOARD)} · 10-day profit "
                         f"<b>{aed(b['total'])}</b> · ordered {b['ordered']}, sold {b['sold']}, unmet {b['unmet']}, "
                         f"spoiled {b['spoiled']}"),
              picker=dict(kind="select", label="Pin a manager", options=list(lookup_b), fmt=str))
        st.caption("Totals are sums of each day's profit rounded to the cent, exactly as shown in the ledgers, so the "
                   "verdict always matches the displayed numbers. The random manager's orders are fixed per challenge. "
                   "Ten days is a small sample: one lucky or unlucky busy day can decide it. The monthly scoreboard "
                   "(tab E) is the fair long-run comparison.")

# ================================================================ G · Experiments (multi-seed)
PRESETS = {"Spoilage cost up (1.5 → 3.0 AED)": dict(spoilage_cost=3.0),
           "Stockout penalty up (6 → 12 AED)": dict(stockout_penalty=12.0),
           "Demand up (mean 7 → 9)": dict(mean_demand=9),
           "Spoil fraction up (25% → 50%)": dict(spoil_frac=0.5)}
THEORY = {"Spoilage cost up (1.5 → 3.0 AED)": ("Over-stocking gets dearer → expect leaner stock: fewer units ordered.", "mean_order", -1),
          "Stockout penalty up (6 → 12 AED)": ("Shortages get dearer → expect a bigger safety buffer: higher service.", "service", +1),
          "Demand up (mean 7 → 9)": ("More customers → expect bigger orders (the 15-unit fridge may start to bind).", "mean_order", +1),
          "Spoil fraction up (25% → 50%)": ("Carry-over is worth less → expect fewer units held overnight (lower leftover).", "mean_leftover", -1)}
PNAME = {"spoilage_cost": ("spoilage cost", lambda v: f"AED {v:.2f}"),
         "stockout_penalty": ("stockout penalty", lambda v: f"AED {v:.2f}"),
         "mean_demand": ("mean demand", lambda v: f"{v:g}/day"), "spoil_frac": ("spoil fraction", lambda v: f"{v:.0%}")}
EXP_METRICS = [("Profit / month", "mean_profit"), ("Service", "service"), ("Spoilage KPI", "spoilage_kpi"),
               ("Stockout rate", "stockout_rate"), ("Units ordered", "mean_order"), ("Leftover units", "mean_leftover"),
               ("Order cost", "mean_order_cost"), ("Holding", "mean_holding_cost"), ("Spoilage cost", "mean_spoilage_cost"),
               ("Stockout penalty", "mean_stockout_penalty")]


def experiment_summary(name: str, seeds: tuple) -> Dict[str, object]:
    """Train baseline + scenario agents on every seed (cached), evaluate ①②③ on the same held-out months."""
    base_items, scen_items = items_of(DEFAULTS), items_of({**DEFAULTS, **PRESETS[name]})
    pb = {s: tuple(int(x) for x in policy_from_Q(cached_train(base_items, episodes, alpha, gamma, s)[0])) for s in seeds}
    ps = {s: tuple(int(x) for x in policy_from_Q(cached_train(scen_items, episodes, alpha, gamma, s)[0])) for s in seeds}
    (pkey, newv), = PRESETS[name].items()
    plabel, pfmt = PNAME[pkey]
    runs = {"①": dict(label="① Baseline policy · baseline economics", setting=f"{plabel} {pfmt(DEFAULTS[pkey])}",
                      ev=[eval_pi(base_items, pb[s]) for s in seeds]),
            "②": dict(label="② Baseline policy · scenario economics (not retrained)", setting=f"{plabel} {pfmt(newv)}",
                      ev=[eval_pi(scen_items, pb[s]) for s in seeds]),
            "③": dict(label="③ Retrained policy · scenario economics", setting=f"{plabel} {pfmt(newv)}",
                      ev=[eval_pi(scen_items, ps[s]) for s in seeds])}
    for rn in runs.values():
        rn["stats"] = {k: dict(mean=float(np.mean([e[k] for e in rn["ev"]])), lo=float(np.min([e[k] for e in rn["ev"]])),
                               hi=float(np.max([e[k] for e in rn["ev"]])), vals=[float(e[k]) for e in rn["ev"]])
                       for _, k in EXP_METRICS}
    theory, tkey, tsign = THEORY[name]
    deltas = [runs["③"]["ev"][i][tkey] - runs["①"]["ev"][i][tkey] for i in range(len(seeds))]
    agree = sum(1 for x in deltas if x * tsign > 0)
    gains = [runs["③"]["ev"][i]["mean_profit"] - runs["②"]["ev"][i]["mean_profit"] for i in range(len(seeds))]
    d_units = runs["③"]["stats"]["mean_order"]["mean"] - runs["①"]["stats"]["mean_order"]["mean"]
    d_serv = runs["③"]["stats"]["service"]["mean"] - runs["①"]["stats"]["service"]["mean"]
    d_spoil = runs["③"]["stats"]["spoilage_kpi"]["mean"] - runs["①"]["stats"]["spoilage_kpi"]["mean"]
    n = len(seeds)
    shift = (f"Across {n} training seeds, the retrained agent orders {abs(d_units):.1f} units/month "
             f"{'more' if d_units >= 0 else 'fewer'} than the baseline agent; service moved {d_serv * 100:+.1f} pts and the "
             f"spoilage KPI {d_spoil * 100:+.1f} pts. Retraining beat keeping the old policy under the new economics in "
             f"{sum(1 for g_ in gains if g_ > 0)} of {n} seeds (mean {saed(float(np.mean(gains)))}/month).")
    verdict = (f"Theory expectation met in {agree} of {n} seeds." if agree else
               f"Theory expectation NOT met in any of {n} seeds — report this honestly.")
    return dict(name=name, seeds=seeds, pb=pb, ps=ps, runs=runs, theory=theory, agree=agree, n=n, shift=shift,
                verdict=verdict, plabel=plabel, base_setting=runs["①"]["setting"], scen_setting=runs["②"]["setting"],
                base_items=base_items, scen_items=scen_items)


with tabs[6]:
    seeds = tuple(seed + i for i in range(N_EXP_SEEDS))
    st.markdown(f"Each experiment changes **one** business setting from the assignment defaults, **retrains** fresh agents "
                f"on **{N_EXP_SEEDS} training seeds** ({', '.join(map(str, seeds))}) with the same hyperparameters, and "
                f"evaluates every policy on the same {EVAL_MONTHS} held-out months. Averaging over seeds separates a real "
                "policy shift from the luck of a single training run. This tab is independent of the sidebar business sliders.")
    name = st.selectbox("Scenario", list(PRESETS), key="exp_preset")
    (pkey_, newv_), = PRESETS[name].items()
    st.markdown(f"**Baseline:** {PNAME[pkey_][0]} {PNAME[pkey_][1](DEFAULTS[pkey_])} (assignment default)  \n"
                f"**Scenario:** {PNAME[pkey_][0]} {PNAME[pkey_][1](newv_)} — everything else unchanged  \n"
                f"**Training:** {episodes:,} episodes, α={alpha}, γ={gamma}, seeds {', '.join(map(str, seeds))} (from sidebar)")
    st.caption("Theory expectation (to check against the result, not to assume): " + THEORY[name][0])
    run_key = (name, episodes, alpha, gamma, seeds)
    if st.button("🔁 Retrain baseline + scenario", key="btn_exp", type="primary"):
        st.session_state.setdefault("exp_done", set()).add(run_key)
    if run_key in st.session_state.get("exp_done", set()):
        prog = st.progress(0.0, text=f"Training {2 * N_EXP_SEEDS} agents (results are cached after the first run)…")
        base_items_, scen_items_ = items_of(DEFAULTS), items_of({**DEFAULTS, **PRESETS[name]})
        jobs = [(it, s) for s in seeds for it in (base_items_, scen_items_)]
        for i, (it, s) in enumerate(jobs):
            cached_train(it, episodes, alpha, gamma, s)
            prog.progress((i + 1) / len(jobs), text=f"Trained {i + 1} of {len(jobs)} agents…")
        prog.empty()
        X = experiment_summary(name, seeds)
        st.session_state.setdefault("exp_store", {})[run_key] = X
        ekey = fingerprint(run_key)
        reach_b = reachable_states(CafeSim(**dict(X["base_items"])).params())
        reach_s = reachable_states(CafeSim(**dict(X["scen_items"])).params())
        row_b = f"Before · baseline ({X['base_setting']})"
        row_s = f"After · retrained ({X['scen_setting']})"
        hrows = []
        for lab, pols, rch, case in [(row_b, X["pb"], reach_b, "Baseline"), (row_s, X["ps"], reach_s, "Scenario")]:
            for s in range(16):
                vals = [pols[sd][s] for sd in seeds]
                mean_o = float(np.mean(vals))
                hrows.append({"_id": f"{case}:{s}", "row": lab, "case": case, "stock": s, "order": mean_o,
                              "per_seed": " · ".join(str(v_) for v_ in vals), "spread": max(vals) - min(vals),
                              "reach": "reachable" if s in rch else "unreachable (not learned)",
                              "txt": (f"{mean_o:.0f}" if mean_o == int(mean_o) else f"{mean_o:.1f}") + ("" if s in rch else "*")})
        hdf = pd.DataFrame(hrows)
        lookup_hm = {r_["_id"]: r_ for r_ in hrows}
        htips = [T("row:N", title="Policy / setting"), T("stock:O", title="Stock on hand (units)"),
                 T("order:Q", title=f"Mean recommended order ({N_EXP_SEEDS} seeds)", format=".2f"),
                 T("per_seed:N", title="Order on each seed"), T("spread:Q", title="Max − min across seeds"),
                 T("reach:N", title="State")]
        hrect = alt.Chart(hdf).mark_rect(stroke=CARD, strokeWidth=2, cornerRadius=3).encode(
            x=alt.X("stock:O", title="Stock on hand (state)   * = unreachable under that setting", axis=alt.Axis(labelAngle=0)),
            y=alt.Y("row:N", title=None, sort=[row_b, row_s], axis=alt.Axis(labelLimit=320)),
            color=alt.Color("order:Q", title=["Mean order", "(units)"], scale=alt.Scale(scheme="cividis", domain=[0, 15]),
                            legend=alt.Legend(orient="right", gradientLength=90)),
            tooltip=htips).add_params(pick_param("pick_exp_pol"))
        htxt = alt.Chart(hdf).mark_text(fontWeight="bold", fontSize=12).encode(
            x="stock:O", y=alt.Y("row:N", sort=[row_b, row_s]), text="txt:N", tooltip=htips,
            color=alt.condition(alt.datum.order < 9, alt.value("white"), alt.value(NAVY)))
        panel(styled((hrect + htxt).properties(height=230), "Policy before vs after retraining",
                     [f"Mean order per stock level over {N_EXP_SEEDS} training seeds", HINT]),
              f"ch_exp_pol_{ekey}", ["pick_exp_pol"], lookup_hm,
              lambda r_: (f"<b>{r_['row']}</b> · stock <b>{r_['stock']}</b> → mean order <b>{r_['order']:.2f} units</b> "
                          f"(per seed: {r_['per_seed']}; spread {r_['spread']}) · {r_['reach']}"),
              picker=dict(kind="select", label="Pin a cell", options=list(lookup_hm),
                          fmt=lambda i: f"{'Before' if i.startswith('Baseline') else 'After'} · stock {i.split(':')[1]}"))
        trows = []
        for rk in ("①", "②", "③"):
            rn = X["runs"][rk]
            row = {"Run": rn["label"], "Economics": rn["setting"]}
            for lab, k in EXP_METRICS:
                s_ = rn["stats"][k]
                row[lab] = s_["mean"]
                if k in ("mean_profit", "service"):
                    row[f"{lab} range"] = (f"{s_['lo']:,.2f} to {s_['hi']:,.2f}" if k == "mean_profit"
                                           else f"{s_['lo']:.1%} to {s_['hi']:.1%}")
            trows.append(row)
        tdf = pd.DataFrame(trows)
        st.dataframe(tdf[["Run", "Economics", "Profit / month", "Profit / month range", "Service", "Service range",
                          "Spoilage KPI", "Stockout rate", "Units ordered", "Leftover units", "Order cost", "Holding",
                          "Spoilage cost", "Stockout penalty"]], hide_index=True, column_config={
            **{c: st.column_config.NumberColumn(format="%.2f") for c in
               ["Profit / month", "Order cost", "Holding", "Spoilage cost", "Stockout penalty"]},
            **{c: st.column_config.NumberColumn(format="%.3f") for c in ["Service", "Spoilage KPI", "Stockout rate"]},
            **{c: st.column_config.NumberColumn(format="%.1f") for c in ["Units ordered", "Leftover units"]}})
        st.caption(f"Values are means over {N_EXP_SEEDS} training seeds; ranges show the lowest and highest seed.")
        erows, eid = [], {}
        for rk in ("①", "②", "③"):
            rn = X["runs"][rk]
            s_ = rn["stats"]["mean_profit"]
            r_ = dict(_id=rk, run=rn["label"], setting=rn["setting"], profit=s_["mean"], lo=s_["lo"], hi=s_["hi"],
                      service=rn["stats"]["service"]["mean"], spoil=rn["stats"]["spoilage_kpi"]["mean"],
                      stockout=rn["stats"]["stockout_rate"]["mean"], units=rn["stats"]["mean_order"]["mean"],
                      label=aed(s_["mean"]), per_seed=" · ".join(aed(v_) for v_ in s_["vals"]))
            erows.append(r_)
            eid[rk] = r_
            for sd, v_ in zip(seeds, s_["vals"]):
                eid[f"{rk}|{sd}"] = dict(_id=f"{rk}|{sd}", run=rn["label"], setting=rn["setting"], seed=sd, profit=v_,
                                         single=True)
        edf = pd.DataFrame(erows)
        sdots = pd.DataFrame([dict(_id=k, run=v_["run"], setting=v_["setting"], seed=v_["seed"], profit=v_["profit"])
                              for k, v_ in eid.items() if v_.get("single")])
        etips = [T("run:N", title="Run"), T("setting:N", title="Economics"),
                 T("profit:Q", title=f"Mean monthly profit, {N_EXP_SEEDS} seeds (AED)", format=",.2f"),
                 T("per_seed:N", title="Each seed (AED)"), T("service:Q", title="Service", format=".1%"),
                 T("spoil:Q", title="Spoilage KPI", format=".1%"), T("stockout:Q", title="Stockout rate", format=".1%"),
                 T("units:Q", title="Units ordered / month", format=",.1f")]
        eyy = alt.Y("run:N", title=None, sort=list(edf["run"]), axis=alt.Axis(labelLimit=380))
        elo, ehi = min(edf["lo"].min(), 0), max(edf["hi"].max(), 0)
        espan = ehi - elo or 1
        ex = alt.X("profit:Q", title="Mean monthly profit (AED) · whisker = lowest–highest seed",
                   scale=alt.Scale(domain=[elo - 0.04 * espan, ehi + 0.28 * espan]))
        ebar = alt.Chart(edf).mark_bar(cornerRadiusEnd=3, height=24, opacity=.9).encode(
            x=ex, y=eyy, color=alt.Color("run:N", legend=None, scale=alt.Scale(range=[SLATE, AMBER, EMERALD])),
            tooltip=etips).add_params(pick_param("pick_exp_bar"))
        ewhisk = alt.Chart(edf).mark_rule(color=NAVY, size=1.6).encode(x="lo:Q", x2="hi:Q", y=eyy)
        edots = alt.Chart(sdots).mark_point(color=NAVY, filled=True, size=45, opacity=.85).encode(
            x="profit:Q", y=alt.Y("run:N", sort=list(edf["run"])),
            tooltip=[T("run:N", title="Run"), T("seed:Q", title="Training seed"),
                     T("profit:Q", title="Monthly profit (AED)", format=",.2f")]).add_params(pick_param("pick_exp_dot"))
        etext = alt.Chart(edf).mark_text(align="left", dx=8, color=NAVY, fontWeight="bold").encode(
            x="hi:Q", y=eyy, text="label:N", tooltip=etips)

        def describe_e(r_: dict) -> str:
            if r_.get("single"):
                return f"<b>{r_['run']}</b> ({r_['setting']}) · training seed <b>{r_['seed']}</b> · monthly profit <b>{aed(r_['profit'])}</b>"
            return (f"<b>{r_['run']}</b> ({r_['setting']}) · mean monthly profit <b>{aed(r_['profit'])}</b> "
                    f"(seeds: {r_['per_seed']}) · service {pct(r_['service'])} · spoilage KPI {pct(r_['spoil'])} · "
                    f"stockout {pct(r_['stockout'])} · {r_['units']:,.1f} units ordered/month")

        panel(styled((ebar + ewhisk + edots + etext).properties(height=190), "What-if: profit per run",
                     [f"Bars = mean of {N_EXP_SEEDS} seeds · dots = individual seeds", HINT]),
              f"ch_exp_bar_{ekey}", ["pick_exp_bar", "pick_exp_dot"], eid, describe_e,
              picker=dict(kind="select", label="Pin a run or seed", options=list(eid),
                          fmt=lambda i: (i if "|" not in i else f"{i.split('|')[0]} seed {i.split('|')[1]}")))
        (st.success if X["agree"] == X["n"] else st.warning if X["agree"] else st.error)(
            f"**Observed shift (computed).** {X['shift']} **{X['verdict']}**")
    else:
        st.info(f"Press **Retrain baseline + scenario**. The first run trains {2 * N_EXP_SEEDS} agents (about 1–1.5 "
                "minutes); later views are instant because results are cached. Baseline agents are shared between "
                "scenarios, so the next scenario needs only 3 new trainings.")


# ================================================================ H · Assignment (+ Word evidence pack)
def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 8.5, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#9aa3ad", "axes.labelcolor": NAVY, "axes.titlecolor": NAVY,
                         "axes.titleweight": "bold", "axes.titlesize": 10, "axes.titlelocation": "left",
                         "xtick.color": SLATE, "ytick.color": SLATE, "axes.grid": True, "grid.color": "#e7e1d5",
                         "grid.linewidth": .7, "figure.facecolor": "white", "axes.facecolor": "white",
                         "legend.frameon": False, "axes.axisbelow": True})
    return plt


def _png(plt, fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", dpi=200)
    plt.close(fig)
    buf.seek(0)
    return buf


def build_evidence_pack(ctx: dict) -> bytes:
    """Word evidence pack built from THIS run's numbers (no hard-coded results)."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    plt = _plt()
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())  # noqa: E731
    strip_md = lambda s: re.sub(r"\*\*(.+?)\*\*", r"\1", s)         # noqa: E731
    p_, res, const, best_k = ctx["P"], ctx["res"], ctx["const"], ctx["best_k"]
    r = res["RL agent (greedy)"]
    cfg = ctx["cfg"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.font.color.rgb = rgb("#1F2A37")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for lvl, size in ((1, 15), (2, 12)):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name, hs.font.size, hs.font.bold = "Calibri", Pt(size), True
        hs.font.color.rgb = rgb(NAVY)
        hs.paragraph_format.space_before, hs.paragraph_format.space_after = Pt(14 if lvl == 1 else 8), Pt(4)
        hs.paragraph_format.keep_with_next = True
        rpr = hs.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Calibri")

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill.lstrip("#"))
        tcPr.append(shd)

    def borders(table, color="D9D2C3"):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            b.append(el)
        tblPr.append(b)

    def rule(par, color=TEAL, sz=14):
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single")
        bt.set(qn("w:sz"), str(sz))
        bt.set(qn("w:space"), "4")
        bt.set(qn("w:color"), color.lstrip("#"))
        bdr.append(bt)
        pPr.append(bdr)

    def fix_widths(t, widths):
        """Word and LibreOffice read widths from the grid AND the cells; fixed layout stops auto-resizing."""
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))
        tw.set(qn("w:type"), "dxa")
        for old in t._tbl.tblPr.findall(qn("w:tblW")):
            t._tbl.tblPr.remove(old)
        t._tbl.tblPr.append(tw)
        for i, wcm in enumerate(widths):
            t.columns[i].width = Cm(wcm)
        for row in t.rows:
            for i, wcm in enumerate(widths):
                row.cells[i].width = Cm(wcm)

    def table(header, rows, widths, align=None, font=9, bold_rows=()):
        t = doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            run_ = par.add_run(h)
            run_.bold, run_.font.size = True, Pt(font)
            run_.font.color.rgb = rgb("#FFFFFF")
            if align and align[i] == "r":
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            shade(c, NAVY)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                run_ = par.add_run(str(v))
                run_.font.size = Pt(font)
                run_.bold = ri in bold_rows
                if align and align[i] == "r":
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if ri % 2 == 1:
                    shade(cells[i], "F6F3EC")
        fix_widths(t, widths)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def callout(text, fill="EEF6F4", edge=TEAL, size=9.5, italic=False):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t, color=edge.lstrip("#"))
        c = t.rows[0].cells[0]
        fix_widths(t, [17.0])
        no_split = OxmlElement("w:cantSplit")          # keep each callout box on one page
        no_split.set(qn("w:val"), "true")
        t.rows[0]._tr.get_or_add_trPr().append(no_split)
        shade(c, fill)
        par = c.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        run_ = par.add_run(text)
        run_.font.size, run_.italic = Pt(size), italic
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(buf, caption):
        doc.add_picture(buf, width=Cm(16.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        run_ = cap.add_run(caption)
        run_.italic, run_.font.size = True, Pt(8.5)
        run_.font.color.rgb = rgb(SLATE)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def para(text, size=None, color=None, bold=False, italic=False, after=None):
        par = doc.add_paragraph()
        run_ = par.add_run(text)
        if size:
            run_.font.size = Pt(size)
        if color:
            run_.font.color.rgb = rgb(color)
        run_.bold, run_.italic = bold, italic
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        return par

    def bullet(text):
        par = doc.add_paragraph(style="List Bullet")
        par.add_run(text)
        par.paragraph_format.space_after = Pt(2)

    # ---- footer with page numbers
    fpar = sec.footer.paragraphs[0]
    fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fpar.add_run("The Smart Café · Q-learning evidence pack · simulated data · page ")
    frun.font.size, frun.font.color.rgb = Pt(8), rgb(SLATE)
    prun = fpar.add_run()
    prun.font.size, prun.font.color.rgb = Pt(8), rgb(SLATE)
    for tag, txt in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        prun._r.append(el)

    # ---- cover block
    para("THE SMART CAFÉ · REINFORCEMENT LEARNING LAB", size=9, color=TEAL, bold=True, after=2)
    para("Q-Learning Evidence Pack", size=24, color=NAVY, bold=True, after=0)
    t_ = para("Learning to order a perishable item from profit alone", size=12, color=SLATE, after=6)
    rule(t_)
    env_txt = (f"spoilage AED {p_['spoilage_cost']:.2f}, stockout penalty AED {p_['stockout_penalty']:.2f}, "
               f"mean demand {p_['mean_demand']:g}/day (+20% busy days), {p_['spoil_frac']:.0%} of leftovers spoil; "
               f"price AED {p_['price']:.2f}, order AED {p_['order_cost']:.2f}, holding AED {p_['holding_cost']:.2f}, "
               f"fridge {p_['max_stock']}, start stock {p_['start_stock']}, {p_['days']}-day months")
    table(["Item", "Details"], [
        ["Prepared by", ctx["name"] or "______________________"],
        ["Generated", ctx["when"]],
        ["Training", f"{cfg['episodes']:,} episodes · α {cfg['alpha']} · γ {cfg['gamma']} · ε 1.0 → 0.05 (×0.9998 per episode) · seed {cfg['seed']}"],
        ["Business settings", env_txt],
        ["Evaluation", f"{EVAL_MONTHS} held-out months (seed {EVAL_SEED}); every policy sees identical demand"]],
        [3.6, 13.4])
    callout("All figures are simulated with the course simulator; they are not real café sales. Numbers in this pack "
            "come from the dashboard run described above and change if the settings or seed change.")

    # ---- 1. headline
    doc.add_heading("1. Headline results", level=1)
    gap = r["mean_profit"] - const[best_k]["mean_profit"]
    table(["Metric", "Value", "Definition"], [
        ["Average monthly net profit", aed(r["mean_profit"]), f"mean over {EVAL_MONTHS} held-out months (std {aed(r['std_profit'])})"],
        ["Service level", pct(r["service"]), "units sold ÷ units demanded"],
        ["Spoilage KPI", pct(r["spoilage_kpi"]), "units spoiled ÷ units ordered"],
        ["Stockout rate", pct(r["stockout_rate"]), "unmet units ÷ units demanded"],
        ["Net margin", pct(r["net_margin"]), "net profit ÷ revenue"],
        ["Loss-making months", pct(r["loss_share"]), "share of months with profit < 0"],
        ["vs best fixed order", saed(gap), f"best constant order = {best_k}/day ({aed(const[best_k]['mean_profit'])})"]],
        [5.0, 3.2, 8.8], align="lrl")
    para(f"The RL agent's average month: revenue {aed(r['mean_revenue'])} minus order cost {aed(r['mean_order_cost'])}, "
         f"holding {aed(r['mean_holding_cost'])}, spoilage {aed(r['mean_spoilage_cost'])} and stockout penalty "
         f"{aed(r['mean_stockout_penalty'])} = net profit {aed(r['mean_profit'])}.")

    # ---- 2. MDP
    doc.add_heading("2. MDP framing (P1)", level=1)
    table(["RL term", "In the café"], [
        ["State", "Stock on hand each morning (0–15 units)"],
        ["Action", "Units ordered today (0–15)"],
        ["Reward", "Today's profit = sales − order cost − holding − spoilage − stockout penalty"],
        ["Episode", "One 30-day month starting with 5 units"],
        ["Policy", "The stock → order rule read from the 16 × 16 Q-table (argmax)"]], [3.2, 13.8])

    # ---- 3. baselines
    doc.add_heading("3. Naive managers vs the agent (P2, P3)", level=1)
    rows_ = []
    rl = r["mean_profit"]
    allp = list(res.items()) + [(f"Best constant order = {best_k} (diagnostic)", const[best_k])]
    allp.sort(key=lambda kv: -kv[1]["mean_profit"])
    for k, v in allp:
        rows_.append([k, aed(v["mean_profit"]), aed(v["std_profit"]), pct(v["service"]), pct(v["stockout_rate"]),
                      pct(v["spoilage_kpi"]), saed(v["mean_profit"] - rl)])
    table(["Manager", "Profit / month", "Std", "Service", "Stockout", "Spoilage KPI", "Δ vs RL"], rows_,
          [4.5, 2.5, 2.3, 1.6, 1.7, 2.0, 2.4], align="lrrrrrr",
          bold_rows=tuple(i for i, (k, _) in enumerate(allp) if k.startswith("RL")))
    fig, ax = plt.subplots(figsize=(6.6, 2.7))
    names_ = [k for k, _ in allp][::-1]
    vals_ = [v["mean_profit"] for _, v in allp][::-1]
    cols_ = [EMERALD if n_.startswith("RL") else BLUE if n_.startswith("Best") else (SLATE if v_ >= 0 else CORAL)
             for n_, v_ in zip(names_, vals_)]
    ax.barh(names_, vals_, color=cols_, height=.62)
    span_ = max(vals_) - min(min(vals_), 0)
    for i, v_ in enumerate(vals_):
        ax.text(v_ + (.01 if v_ >= 0 else -.01) * span_, i, aed(v_), va="center", ha="left" if v_ >= 0 else "right", fontsize=7.5)
    ax.axvline(0, color=SLATE, lw=.8)
    ax.set_xlim(min(min(vals_), 0) - .22 * span_, max(vals_) + .22 * span_)
    ax.set_xlabel("Average monthly profit (AED)")
    ax.set_title(f"Every manager on the same {EVAL_MONTHS} held-out months")
    figure(_png(plt, fig), "Figure 1 · Monthly scoreboard (green = RL agent, blue = diagnostic benchmark, red = loss).")
    fig, ax = plt.subplots(figsize=(6.6, 2.5))
    ks = list(const)
    ax.plot(ks, [const[k]["mean_profit"] for k in ks], marker="s", ms=4, color=BLUE, label="Constant daily order")
    ax.axhline(rl, color=EMERALD, ls="--", lw=1.6, label=f"RL agent ({aed(rl)})")
    ax.axhline(0, color=SLATE, lw=.8)
    ax.set_xticks(ks)
    ax.set_xlabel("Constant order every day (units)")
    ax.set_ylabel("AED / month")
    ax.set_title("Profit of every constant order (diagnostic)")
    ax.legend(loc="lower right")
    figure(_png(plt, fig), "Figure 2 · Brute-force sweep run after training; the agent never saw these numbers.")
    para(f"Two failure modes (P2): ordering 1 unit a day earns {aed(const[1]['mean_profit'])}/month (stockout rate "
         f"{pct(const[1]['stockout_rate'])}); ordering 15 a day earns {aed(const[15]['mean_profit'])}/month (spoilage KPI "
         f"{pct(const[15]['spoilage_kpi'])}). A lost sale costs about {aed(ctx['cu'])} (margin + penalty) against about "
         f"{aed(ctx['co'])} for a wasted unit, so under-ordering is the more expensive mistake.")

    # ---- 4. learning
    doc.add_heading("4. How the agent learned (P4)", level=1)
    c_, eps_ = ctx["curve"], ctx["diag"]["eps"]
    w_ = max(10, min(200, len(c_) // 20))
    roll_ = np.convolve(c_, np.ones(w_) / w_, mode="valid")
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(6.6, 3.6), sharex=True, gridspec_kw={"height_ratios": [2.2, 1]})
    a1.plot(c_, color=BLUE, alpha=.18, lw=.4, label="Raw episode profit")
    a1.plot(np.arange(w_ - 1, len(c_)), roll_, color=NAVY, lw=1.8, label=f"Rolling mean ({w_} episodes)")
    a1.set_ylabel("Training profit (AED)")
    a1.set_title("Learning curve")
    a1.legend(loc="lower right")
    a2.plot(eps_, color=BLUE, lw=1.4)
    a2.axhline(0.05, color=AMBER, ls="--", lw=1.2)
    a2.set_ylabel("ε")
    a2.set_xlabel("Episode")
    figure(_png(plt, fig), "Figure 3 · Training profit per episode (raw and rolling mean) and the exploration rate ε.")
    n_ = max(1, len(c_) // 10)
    para(f"Average training profit rose from {aed(np.mean(c_[:n_]))} in the first {n_:,} episodes to {aed(np.mean(c_[-n_:]))} "
         f"in the last {n_:,}. Training profit includes random exploration, so it sits below the held-out evaluation. "
         f"ε was {epsilon_at(12000):.4f} after 12,000 episodes and first reaches its 0.05 floor at episode "
         f"{floor_episode():,} (0-based).")
    table(["Hyperparameter", "Value", "Role"], [
        ["α (learning rate)", cfg["alpha"], "how far each update moves Q toward the new target"],
        ["γ (discount)", cfg["gamma"], "how much tomorrow's value counts today"],
        ["ε start → floor", "1.0 → 0.05", "probability of taking the random (explore) branch"],
        ["ε decay", "×0.9998 per episode", "moves behaviour from exploring to exploiting"],
        ["Episodes · seed", f"{cfg['episodes']:,} · {cfg['seed']}", "amount of practice · reproducibility"]],
        [4.0, 3.6, 9.4])

    # ---- 5. strategy
    doc.add_heading("5. The strategy the agent discovered (P5)", level=1)
    summ_ = ctx["summary"]
    if summ_["text"]:
        para(strip_md(summ_["text"]))
    pi_, vis_, reach_ = ctx["PI"], ctx["visits"], ctx["reach"]
    prow_ = []
    for s in range(16):
        ev_ = "unreachable" if s not in reach_ else ("under-visited" if vis_[s] < MIN_VISITS else "well visited")
        prow_.append([s, pi_[s], min(s + pi_[s], 15), f"{int(vis_[s]):,}", ev_])
    table(["Stock", "Order", "Shelf after ordering", "Training visits", "Evidence"], prow_, [2.2, 2.2, 4.0, 3.4, 5.2],
          align="rrrrl", font=8.5)
    fig, ax = plt.subplots(figsize=(6.6, 2.6))
    xs = np.arange(16)
    ax.bar(xs, pi_, color=BLUE, alpha=.45, label="Recommended order")
    shelf_ = [min(s + pi_[s], 15) for s in range(16)]
    ax.plot(xs, shelf_, color=NAVY, lw=1.8, marker="o", ms=4, label="Shelf after ordering")
    if summ_["target"] is not None:
        ax.axhline(summ_["target"], color=TEAL, ls="--", lw=1.4, label=f"Typical order-up-to level ≈ {summ_['target']}")
    ax.set_xticks(xs)
    ax.set_ylim(0, 19)
    ax.set_xlabel("Stock on hand this morning (units)")
    ax.set_ylabel("Units")
    ax.set_title("Order-up-to view of the learned policy")
    ax.legend(loc="upper right", ncol=3, fontsize=7.5)
    figure(_png(plt, fig), "Figure 4 · Bars: order per stock level. Line: shelf after ordering. Stocks "
           f"{[s for s in range(16) if s not in reach_]} cannot occur under these settings, so their orders are defaults.")

    # ---- 6. You vs AI
    doc.add_heading("6. You vs the AI (P6)", level=1)
    G = ctx["game"]
    if G:
        R_ = G["result"]
        who = {"you": f"You won by {aed(R_['margin'])}", "ai": f"The AI won by {aed(R_['margin'])}",
               "tie": "Draw — identical profit to the cent"}[R_["winner"]]
        callout(f"{who}. Your 10-day profit {aed(R_['you_total'])} · AI {aed(R_['ai_total'])} · You − AI "
                f"{saed(R_['you_minus_ai'])} (positive = you won).", fill="FBF5E8", edge=AMBER)
        table(["Day", "Demand", "Your order", "Your profit", "AI order", "AI profit"],
              [[a["day"], a["demand"], y["order"], aed(y["profit"]), a["order"], aed(a["profit"])]
               for a, y in zip(R_["ai_rows"], G["rows"])], [1.6, 2.2, 2.6, 3.4, 2.6, 3.4], align="rrrrrr", font=8.5)
        table(["Rank", "Manager (same 10 days)", "10-day profit", "Sold", "Unmet", "Spoiled"],
              [[b["rank"], b["name"], aed(b["total"]), b["sold"], b["unmet"], b["spoiled"]] for b in G["board"]],
              [1.4, 7.0, 3.0, 1.8, 1.8, 2.0], align="rlrrrr", font=8.5,
              bold_rows=tuple(i for i, b in enumerate(G["board"]) if b["name"] == "You"))
        para("Reflect in your own words: where did your guesses differ from the agent's, and on which days?",
             italic=True, color=SLATE, size=9.5)
    else:
        callout("The 10-day challenge was not completed in this session. Play it in tab F and rebuild this pack to "
                "include your result.", fill="FBF5E8", edge=AMBER, italic=True)

    # ---- 7. experiments
    doc.add_heading("7. Experiments: what you reward is what you get (P7)", level=1)
    if ctx["experiments"]:
        for X_ in ctx["experiments"]:
            doc.add_heading(X_["name"], level=2)
            rows_ = []
            for rk in ("①", "②", "③"):
                rn = X_["runs"][rk]
                st_ = rn["stats"]
                rows_.append([rn["label"], aed(st_["mean_profit"]["mean"]),
                              f"{aed(st_['mean_profit']['lo'])} to {aed(st_['mean_profit']['hi'])}",
                              pct(st_["service"]["mean"]), pct(st_["spoilage_kpi"]["mean"]), f"{st_['mean_order']['mean']:.1f}"])
            table(["Run", "Mean profit", "Seed range", "Service", "Spoilage KPI", "Units / month"], rows_,
                  [6.2, 2.4, 3.6, 1.6, 1.8, 1.4], align="lrrrrr", font=8.5)
            para(X_["shift"] + " " + X_["verdict"])
            para("Theory expectation: " + X_["theory"], italic=True, color=SLATE, size=9.5)
    else:
        callout("No experiment was run with these training settings. Run one in tab G and rebuild the pack.",
                fill="FBF5E8", edge=AMBER, italic=True)

    # ---- 8. draft recommendation
    doc.add_heading("8. Draft recommendation to the café owner", level=1)
    cautious = res["Cautious (order 4)"]["mean_profit"]
    gap_txt = ("about the same as" if abs(gap) <= 0.01 * max(abs(const[best_k]["mean_profit"]), 1)
               else f"{aed(abs(gap))} {'more' if gap > 0 else 'less'} than")
    if ctx["experiments"]:
        X0 = ctx["experiments"][0]
        d0 = X0["runs"]["③"]["stats"]["mean_order"]["mean"] - X0["runs"]["①"]["stats"]["mean_order"]["mean"]
        exp_txt = (f" When the {X0['plabel']} changed ({X0['base_setting'][len(X0['plabel']) + 1:]} → "
                   f"{X0['scen_setting'][len(X0['plabel']) + 1:]}), the retrained agent ordered {abs(d0):.1f} units/month "
                   f"{'more' if d0 >= 0 else 'fewer'} on average across {X0['n']} seeds, so retrain whenever costs or demand change.")
    else:
        exp_txt = " If costs or demand change, retrain the agent — the rule shifts with the rewards."
    target_txt = f"about {summ_['target']} units" if summ_["target"] is not None else "the learned level"
    rec = (f"Use the learned ordering rule: each morning, top the shelf up to {target_txt}, ordering more when stock "
           f"is low and little or nothing when it is high. On {EVAL_MONTHS} identical simulated months it earned "
           f"{aed(rl)} per month — {gap_txt} the best fixed order found afterwards ({best_k}/day) and "
           f"{aed(rl - cautious)} more than the cautious manager ({aed(cautious)}) — while serving {pct(r['service'])} of "
           f"demand. A lost sale costs about {aed(ctx['cu'])} against about {aed(ctx['co'])} for a wasted unit, so err on "
           f"the generous side.{exp_txt}")
    callout(rec, fill="EEF6F4", edge=TEAL, size=10)
    para("Draft generated from this run's numbers — rewrite it in your own words before submitting.",
         italic=True, color=SLATE, size=9)

    # ---- 9. evidence index
    doc.add_heading("9. Where to find evidence for P1–P7", level=1)
    table(["Prompt", "Evidence", "Key figure from this run"], [
        ["P1", "Section 2", "state = stock 0–15 · action = order 0–15"],
        ["P2", "Figure 2", f"order 1/day {aed(const[1]['mean_profit'])} · order 15/day {aed(const[15]['mean_profit'])}"],
        ["P3", "Section 3, Figure 1", f"cautious {aed(cautious)} · RL {aed(rl)}"],
        ["P4", "Section 4, Figure 3", f"training profit {aed(np.mean(c_[:n_]))} → {aed(np.mean(c_[-n_:]))}"],
        ["P5", "Section 5, Figure 4", f"order-up-to ≈ {summ_['target']} units" if summ_["target"] is not None else "see table"],
        ["P6", "Section 6", (G and {"you": "you won", "ai": "AI won", "tie": "draw"}[G["result"]["winner"]]) or "not played yet"],
        ["P7", "Section 7", f"{len(ctx['experiments'])} experiment(s) included"]], [1.6, 4.2, 11.2])

    # ---- 10. method
    doc.add_heading("10. Method and limitations", level=1)
    for b_ in [f"Evaluation uses common random numbers: every policy plays the same {EVAL_MONTHS} demand months.",
               "Training reward (Figure 3) includes exploration; all comparisons use the greedy policy on held-out months.",
               "Experiments average 3 training seeds; a single seed is a demonstration, not a settled result.",
               "The state is stock only, so the agent cannot plan for the end of the month.",
               "Unreachable stock levels keep default actions; they are labelled, not interpreted."]:
        bullet(b_)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


with tabs[7]:
    st.markdown("#### MDP framing (Part 1)")
    st.dataframe(pd.DataFrame({
        "RL term": ["State", "Action", "Reward", "Episode", "Policy"],
        "In the café": ["The stock on hand each morning (0–15 units)",
                        "How many units to order today (0–15)",
                        "Today's profit = sales − order cost − holding − spoilage − stockout penalty",
                        "One full month of 30 days, starting with 5 units",
                        "The rule mapping each stock level to an order, read from the Q-table (argmax)"]}),
        hide_index=True)
    st.markdown("#### Reflection prompts and where the evidence is")
    st.dataframe(pd.DataFrame({
        "Prompt": ["P1", "P2", "P3", "P4", "P5", "P6", "P7"],
        "Question (from the guide)": [
            "Fill in state, action, reward, episode, policy.",
            "Order 15 every day vs 1 every day: what happens to spoilage, stockouts and profit?",
            "Which naive manager is best? Does any lose money? Why is too little worse than too much?",
            "What do α, γ and ε control? Why does ε start high and shrink?",
            "Describe the strategy the AI discovered. Does 'order more when stock is low' appear?",
            "Did the AI beat you? Where did your guesses go wrong? Did 'cautious' surprise you?",
            "Pick one experiment: how did the strategy change, and why?"],
        "Evidence in this app": ["Table above", "Tab E constant-order sweep (orders 1 and 15); tab D ledger",
                                 "Tab E scoreboard", "Tab B learning curve + ε chart; tab I",
                                 "Tab C heatmap + order-up-to view", "Tab F verdict + 10-day scoreboard",
                                 "Tab G multi-seed before/after"]}), hide_index=True)

    st.markdown("#### 📄 Evidence pack (Word)")
    st.caption("A formatted Word document built from this run: headline KPIs, scoreboard and charts, learning curve, "
               "learned policy, your You-vs-AI result and 10-day ranking, every experiment you ran with the current "
               "training settings, a draft owner recommendation and an evidence index for P1–P7.")
    if FRESH:
        exps = [v for k, v in st.session_state.get("exp_store", {}).items()
                if k[1:4] == (episodes, alpha, gamma) and k[4] == tuple(seed + i for i in range(N_EXP_SEEDS))]
        ep_name = st.text_input("Your name for the cover (optional)", key="ep_name", max_chars=60)
        pack_key = (fingerprint(CFG), ep_name, GAME["round"], len(GAME["rows"]),
                    tuple(sorted(x["name"] for x in exps)))
        cA, cB = st.columns([1, 2])
        if cA.button("🛠️ Build evidence pack", key="btn_pack", type="primary"):
            try:
                ctx = dict(name=ep_name.strip(), when=dt.datetime.now().strftime("%d %b %Y, %H:%M"), P=P, cfg=CFG,
                           res=SB[0], const=SB[1], best_k=BEST_K, curve=MODEL["curve"], diag=MODEL["diag"], PI=PI,
                           visits=MODEL["diag"]["state_visits"], reach=reachable_states(P),
                           summary=policy_summary(PI, MODEL["diag"]["state_visits"], reachable_states(P)),
                           cu=cu_cost, co=co_cost, experiments=exps,
                           game=(dict(result=GAME_RES, rows=GAME["rows"], board=BOARD) if GAME_RES else None))
                with st.spinner("Building the Word document…"):
                    st.session_state["pack"] = dict(key=pack_key, data=build_evidence_pack(ctx))
            except ImportError:
                st.error("The evidence pack needs **python-docx** and **matplotlib**: run `pip install python-docx matplotlib`.")
        pack = st.session_state.get("pack")
        if pack and pack["key"] == pack_key:
            cB.download_button("⬇️ Download evidence pack (.docx)", pack["data"], "smart_cafe_evidence_pack.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="dl_pack")
        elif pack:
            cB.caption("The run, game, experiments or name changed since the last build — press **Build** again.")
        st.caption(f"Included now: {'your completed challenge' if GAME_RES else 'no completed challenge yet'} · "
                   f"{len(exps)} experiment(s) run with the current training settings.")
    else:
        st.info("Train the agent (on the current settings) to build the evidence pack.")

    st.markdown("#### Data exports (CSV)")
    if FRESH:
        res, const = SB
        buf = io.StringIO()
        wr = csv.writer(buf)
        wr.writerow(["policy", "mean_profit", "std", "service", "stockout_rate", "spoilage_kpi"])
        for k, v in list(res.items()) + [(f"constant order {k}", v) for k, v in const.items()]:
            wr.writerow([k, round(v["mean_profit"], 2), round(v["std_profit"], 2), round(v["service"], 4),
                         round(v["stockout_rate"], 4), round(v["spoilage_kpi"], 4)])
        e1, e2, e3 = st.columns(3)
        e1.download_button("⬇️ Scoreboard CSV", buf.getvalue().encode(), "scoreboard.csv", "text/csv", key="dl_score")
        e2.download_button("⬇️ Policy CSV", ("stock,order,training_visits\n" + "\n".join(
            f"{s},{a},{int(MODEL['diag']['state_visits'][s])}" for s, a in enumerate(PI))).encode(),
            "policy.csv", "text/csv", key="dl_policy")
        e3.download_button("⬇️ Learning curve CSV", ("episode,profit,epsilon\n" + "\n".join(
            f"{i},{v:.2f},{MODEL['diag']['eps'][i]:.6f}" for i, v in enumerate(MODEL["curve"]))).encode(),
            "learning_curve.csv", "text/csv", key="dl_curve")
        st.caption("The 30-day ledger CSV is in tab D.")
    else:
        st.info("Train the agent (on the current settings) to enable exports.")
    st.markdown("#### Screenshot checklist")
    for i, item in enumerate(["Learning curve with ε panel (tab B)", "Policy heatmap + order-up-to view (tab C)",
                              "Monthly scoreboard (tab E)", "Finished You-vs-AI result + 10-day scoreboard (tab F)",
                              "One experiment's before/after (tab G)"]):
        st.checkbox(item, key=f"shot_{i}")

# ================================================================ I · ML Diagnostics
with tabs[8]:
    st.markdown("**State** = stock on hand (0–15) · **Action** = order (0–15) · **Q-table** = 16 × 16 · "
                "**Reward** = daily profit · **Episode** = 30 days · update target = r + γ · max Q(s′, ·)")
    with st.expander("Feature audit: what the agent sees vs what the dashboard shows", expanded=False):
        st.dataframe(pd.DataFrame([
            ["Stock on hand (STATE)", "carry-over from yesterday", "the decision input", "simulator", "Yes — only input"],
            ["Utilisation", "stock ÷ 15", "how full the fridge is", "stock", "No (diagnostic)"],
            ["Base-demand coverage", "stock ÷ mean_demand (n/a if 0)", "days of typical demand on hand", "stock, setting", "No (diagnostic)"],
            ["Model-implied expected demand", "mean_demand + busy_chance × 5", "context (8.0 at defaults)", "settings", "No (diagnostic)"],
            ["Expected-demand coverage", "stock ÷ expected demand", "coverage incl. busy days", "stock, settings", "No (diagnostic)"],
            ["Remaining capacity", "15 − stock", "room to order", "stock", "No (diagnostic)"],
            ["Post-order inventory", "min(stock + order, 15)", "order-up-to check", "stock, policy", "No (diagnostic)"],
            ["Today's realised demand", "drawn after ordering", "outcome only", "simulator", "No — would be leakage"]],
            columns=["Feature", "Formula", "Purpose", "Source", "Used by the agent?"]), hide_index=True)
    fe = floor_episode()
    cards([("Epsilon after 12,000 episodes", f"{epsilon_at(12000):.4f}", "ε = 0.9998^12000", BLUE),
           ("Epsilon after 15,000 episodes", f"{epsilon_at(15000):.4f}", "ε floored at 0.05", BLUE),
           ("First floored episode", f"{fe:,}", "0-based index", SLATE)])
    if need_model():
        d = MODEL["diag"]
        win = 500 if len(d["eps"]) >= 2000 else 100
        k = len(d["eps"]) // win
        if k:
            rb = d["random_branch"][: k * win].reshape(k, win).mean(1) / 30
            ng = d["nongreedy"][: k * win].reshape(k, win).mean(1) / 30
            ex = d["eps"][: k * win].reshape(k, win).mean(1)
            S1, S2, S3 = "Logged: random branch taken", "Logged: random AND ≠ greedy action", "Expected share (mean ε in window)"
            xrows = []
            for i in range(k):
                for sname, arr in ((S1, rb), (S2, ng), (S3, ex)):
                    xrows.append({"_id": f"{sname}|{i}", "series": sname, "mid": (i + 0.5) * win,
                                  "window": f"{i * win:,}–{(i + 1) * win - 1:,}", "share": float(arr[i])})
            xdf = pd.DataFrame(xrows)
            lookup_x = {r_["_id"]: r_ for r_ in xrows}
            xtips = [T("series:N", title="Series"), T("window:N", title="Episodes in window"),
                     T("share:Q", title="Share of decision days", format=".2%")]
            xcol = alt.Color("series:N", title=None, scale=alt.Scale(domain=[S1, S2, S3], range=[BLUE, CORAL, SLATE]))
            xdash = alt.StrokeDash("series:N", legend=None, scale=alt.Scale(domain=[S1, S2, S3], range=[[1, 0], [6, 3], [2, 3]]))
            xline = alt.Chart(xdf).mark_line(strokeWidth=2).encode(
                x=alt.X("mid:Q", title=f"Episode (windows of {win})", scale=alt.Scale(domain=[0, k * win], nice=False)),
                y=alt.Y("share:Q", title="Share of decision days", axis=alt.Axis(format="%")), color=xcol, strokeDash=xdash)
            xpts = alt.Chart(xdf).mark_point(filled=True, size=45).encode(
                x="mid:Q", y="share:Q", color=xcol,
                shape=alt.Shape("series:N", legend=None, scale=alt.Scale(domain=[S1, S2, S3], range=["circle", "triangle-up", "diamond"])),
                tooltip=xtips).add_params(pick_param("pick_expl"))
            panel(styled((xline + xpts).properties(height=280), "Exploration vs exploitation",
                         ["Measured during training (logged branch choices)", HINT]),
                  f"ch_expl_{FP}", ["pick_expl"], lookup_x,
                  lambda r_: f"<b>{r_['series']}</b> · episodes {r_['window']} · share of decision days <b>{r_['share']:.2%}</b>",
                  picker=dict(kind="select", label="Pin a window", options=list(lookup_x),
                              fmt=lambda i: f"{i.split('|')[0].replace('Logged: ', '')[:28]} · window {int(i.split('|')[1]) + 1}"))
            phases = np.array_split(np.arange(len(d["eps"])), 3)
            st.dataframe(pd.DataFrame([{"Phase": p_name, "Episodes": f"{ix[0]:,}–{ix[-1]:,}",
                                        "Random-branch days": int(d["random_branch"][ix].sum()),
                                        "Greedy-branch days": int(d["greedy_branch"][ix].sum()),
                                        "Random share": d["random_branch"][ix].sum() / (30 * len(ix)),
                                        "Random but = greedy": int((d["random_branch"][ix] - d["nongreedy"][ix]).sum())}
                                       for p_name, ix in zip(["Early", "Middle", "Late"], phases)]),
                         hide_index=True, column_config={"Random share": st.column_config.NumberColumn(format="%.3f")})
        cur = d["eps"][-1]
        st.markdown(f"**What ε = {cur:.4f} means now:** on each day the agent takes the random branch with probability "
                    f"{cur:.4f}, so it is **expected** to explore on about **{30 * cur:.1f} of 30 days** per month. A "
                    f"random pick lands on the greedy order 1 time in 16, so about {30 * cur * 15 / 16:.1f} of those days "
                    "are expected to be genuinely non-greedy. These are expectations, not counts — the logged counts "
                    "are in the chart and table above.")
        sv = d["state_visits"]
        reach = reachable_states(P)
        vrows = [{"_id": s, "stock": s, "visits": int(sv[s]),
                  "status": ("unreachable" if s not in reach else f"under-visited (< {MIN_VISITS})" if sv[s] < MIN_VISITS
                             else "well visited")} for s in range(16)]
        vdf = pd.DataFrame(vrows)
        lookup_v = {r_["_id"]: r_ for r_ in vrows}
        vtips = [T("stock:O", title="Stock on hand (state)"), T("visits:Q", title="Training visits", format=",d"),
                 T("status:N", title="Evidence")]
        vbar = alt.Chart(vdf).mark_bar(cornerRadiusEnd=2).encode(
            x=alt.X("stock:O", title="Stock on hand (state)", axis=alt.Axis(labelAngle=0)),
            y=alt.Y("visits:Q", title="Training visits (log-like scale)", scale=alt.Scale(type="symlog"),
                    axis=alt.Axis(values=[0, 10, 100, 1000, 10000, 100000], format=",d")),
            color=alt.Color("status:N", title=None, scale=alt.Scale(
                domain=["well visited", f"under-visited (< {MIN_VISITS})", "unreachable"], range=[BLUE, AMBER, "#cfc7b8"])),
            tooltip=vtips).add_params(pick_param("pick_vis"))
        # thin hit-strip on the zero line so zero-visit states can still be hovered and clicked
        vhit = alt.Chart(vdf).mark_tick(thickness=8, size=30, opacity=0.35, color="#cfc7b8").encode(
            x="stock:O", y=alt.datum(0), tooltip=vtips).add_params(pick_param("pick_vis0"))
        panel(styled((vbar + vhit).properties(height=280), "State visitation",
                     ["How often each stock level was seen in training", HINT]),
              f"ch_vis_{FP}", ["pick_vis", "pick_vis0"], lookup_v,
              lambda r_: f"Stock <b>{r_['stock']}</b> · <b>{r_['visits']:,}</b> training visits · {r_['status']}",
              picker=dict(kind="select", label="Pin a stock level", options=list(range(16)), fmt=lambda s: f"Stock {s}"))
        sa = d["sa_visits"]
        tried = int((sa[sorted(reach)] > 0).sum())
        st.caption(f"{tried} of {len(reach) * 16} reachable stock–order pairs were tried at least once. "
                   "No confidence score is shown: visit counts are the honest measure of evidence here.")
        with st.expander("Optional: ε-decay sensitivity (only the decay changes; 3 seeds each; held-out months)"):
            st.markdown(f"Compares decay 0.999, 0.9995 and the assignment default 0.9998 with everything else fixed "
                        f"({episodes:,} episodes, α={alpha}, γ={gamma}, current business settings). Each greedy policy is "
                        f"scored on {EVAL_MONTHS} held-out months (seed {SENS_SEED}). Runs 9 trainings — about 1–2 minutes.")
            if st.button("Run sensitivity study", key="btn_sens"):
                out, per = [], []
                for dec in (0.999, 0.9995, 0.9998):
                    vals = []
                    for sd in (11, 12, 13):
                        v_ = eval_pi(ENV, tuple(int(x) for x in policy_from_Q(
                            cached_train(ENV, episodes, alpha, gamma, sd, dec)[0])), SENS_SEED)["mean_profit"]
                        vals.append(v_)
                        per.append(dict(_id=f"{dec}|{sd}", decay=str(dec), seed=sd, profit=float(v_)))
                    out.append({"ε decay": dec, "Default?": "✓" if dec == 0.9998 else "",
                                "Mean profit / month": float(np.mean(vals)), "Std (3 seeds)": float(np.std(vals, ddof=1)),
                                "Min": float(min(vals)), "Max": float(max(vals)),
                                "Floor reached at episode": floor_episode(eps_decay=dec)})
                st.session_state["sens"] = dict(key=CFG, rows=out, per=per)
            sens = st.session_state.get("sens")
            if sens and sens["key"] == CFG:
                st.dataframe(pd.DataFrame(sens["rows"]), hide_index=True, column_config={
                    c: st.column_config.NumberColumn(format="%.2f") for c in
                    ["Mean profit / month", "Std (3 seeds)", "Min", "Max"]})
                mrows = [dict(_id=f"{r_['ε decay']}|mean", decay=str(r_["ε decay"]), profit=r_["Mean profit / month"],
                              lo=r_["Min"], hi=r_["Max"], std=r_["Std (3 seeds)"],
                              default="assignment default" if r_["Default?"] else "alternative") for r_ in sens["rows"]]
                lookup_sn = {x["_id"]: x for x in sens.get("per", []) + mrows}
                sx = alt.X("decay:N", title="ε decay per episode", sort=[str(x) for x in (0.999, 0.9995, 0.9998)],
                           axis=alt.Axis(labelAngle=0))
                sy = alt.Y("profit:Q", title="Held-out monthly profit (AED)", scale=alt.Scale(zero=False))
                srule = alt.Chart(pd.DataFrame(mrows)).mark_rule(color=SLATE, size=2).encode(x=sx, y="lo:Q", y2="hi:Q")
                smean = alt.Chart(pd.DataFrame(mrows)).mark_tick(color=NAVY, thickness=3, size=34).encode(
                    x=sx, y=sy, tooltip=[T("decay:N", title="ε decay"), T("default:N", title="Schedule"),
                                         T("profit:Q", title="Mean of 3 seeds (AED)", format=",.2f"),
                                         T("std:Q", title="Std across seeds (AED)", format=",.2f")]).add_params(pick_param("pick_sens_m"))
                sdot = alt.Chart(pd.DataFrame(sens.get("per", []))).mark_point(color=BLUE, filled=True, size=80).encode(
                    x=sx, y=sy, tooltip=[T("decay:N", title="ε decay"), T("seed:Q", title="Training seed"),
                                         T("profit:Q", title="Monthly profit (AED)", format=",.2f")]).add_params(pick_param("pick_sens_d"))
                panel(styled((srule + smean + sdot).properties(height=260), "ε-decay sensitivity",
                             ["Dots = seeds · bar = mean · line = range", HINT]),
                      f"ch_sens_{FP}", ["pick_sens_m", "pick_sens_d"], lookup_sn,
                      lambda x: (f"ε decay <b>{x['decay']}</b> · seed <b>{x['seed']}</b> · monthly profit <b>{aed(x['profit'])}</b>"
                                 if "seed" in x else
                                 f"ε decay <b>{x['decay']}</b> ({x['default']}) · mean of 3 seeds <b>{aed(x['profit'])}</b> "
                                 f"(range {aed(x['lo'])} to {aed(x['hi'])}, std {aed(x['std'])})"),
                      picker=dict(kind="select", label="Pin a seed or mean", options=list(lookup_sn),
                                  fmt=lambda i: f"decay {i.split('|')[0]} · " + ("mean" if i.endswith("mean") else f"seed {i.split('|')[1]}")))
                st.caption("Three seeds is a small sample — differences smaller than about one Std are not evidence "
                           "that one schedule is better. The assignment default remains the primary model.")
